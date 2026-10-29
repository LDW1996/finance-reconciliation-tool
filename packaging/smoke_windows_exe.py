from __future__ import annotations

from io import BytesIO
import json
from email.message import EmailMessage
import os
import re
import subprocess
import sys
import tempfile
import time
from pathlib import Path

import httpx
import pandas as pd
from openpyxl import load_workbook


APP_URL = "http://127.0.0.1:8765"


def wait_for_health(timeout_seconds: int = 45, process=None, startup_log=None) -> None:
    deadline = time.time() + timeout_seconds
    last_error: Exception | None = None
    while time.time() < deadline:
        if process is not None and process.poll() is not None:
            detail=Path(startup_log).read_text(encoding='utf-8',errors='replace')[-24000:] if startup_log and Path(startup_log).exists() else 'No startup log was produced.'
            raise RuntimeError(f'EXE exited before health check, code={process.poll()}:\n{detail}')
        try:
            response = httpx.get(f"{APP_URL}/api/health", timeout=2)
            if response.status_code == 200 and response.json().get("status") == "ok":
                return
        except Exception as exc:
            last_error = exc
        time.sleep(1)
    detail=Path(startup_log).read_text(encoding='utf-8',errors='replace')[-24000:] if startup_log and Path(startup_log).exists() else 'No startup log was produced.'
    raise RuntimeError(f"EXE did not expose /api/health in time: {last_error}\n{detail}")


def make_sample_workbook(path: Path) -> None:
    rows = [
        {"公司代码or伙伴公司": "1703", "凭证编号": "A-1", "分配": "OK001", "文本": "抵消", "原币金额": 100},
        {"公司代码or伙伴公司": "2615", "凭证编号": "B-1", "分配": "OK001", "文本": "抵消", "原币金额": -100},
        {"公司代码or伙伴公司": "1703", "凭证编号": "A-2", "分配": "DIFF001", "文本": "差异", "原币金额": 100},
        {"公司代码or伙伴公司": "2615", "凭证编号": "B-2", "分配": "DIFF001", "文本": "差异", "原币金额": 80},
        {"公司代码or伙伴公司": "1703", "凭证编号": "A-3", "分配": "", "文本": "空分配", "原币金额": 50},
    ]
    pd.DataFrame(rows).to_excel(path, index=False)


def assert_frontend_served() -> None:
    response = httpx.get(APP_URL, timeout=10)
    response.raise_for_status()
    if '<div id="app"></div>' not in response.text:
        raise AssertionError("Frontend index.html was not served by the EXE")
    asset = re.search(r'src="(/assets/[^"]+\.js)"', response.text)
    if not asset:
        raise AssertionError("Frontend JavaScript asset was not linked")
    script = httpx.get(f"{APP_URL}{asset.group(1)}", timeout=10)
    script.raise_for_status()
    if "原表标色" not in script.text:
        raise AssertionError("The EXE contains an outdated frontend build")
    commit = os.environ.get("GITHUB_SHA", "")
    if commit and commit[:8] not in script.text:
        raise AssertionError("The EXE does not contain the current workflow commit")


def assert_excel_flow(sample_path: Path) -> None:
    with sample_path.open("rb") as file_obj:
        response = httpx.post(
            f"{APP_URL}/api/inspect-combined",
            files={"workbook": (sample_path.name, file_obj, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            data={"match_field": "分配", "amount_field": "原币金额"},
            timeout=20,
        )
    response.raise_for_status()
    report = response.json()
    if not report.get("valid"):
        raise AssertionError(f"Inspect endpoint rejected valid workbook: {report}")
    if report.get("companyCodes") != ["1703", "2615"]:
        raise AssertionError(f"Unexpected company codes: {report}")

    with sample_path.open("rb") as file_obj:
        response = httpx.post(
            f"{APP_URL}/api/analyze-combined",
            files={"workbook": (sample_path.name, file_obj, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            data={"match_field": "分配", "amount_field": "原币金额"},
            timeout=30,
        )
    response.raise_for_status()
    if response.headers.get("x-allocation-matched") != "1":
        raise AssertionError(f"Unexpected matched count headers: {response.headers}")
    if len(response.content) < 5000:
        raise AssertionError("Generated Excel response is unexpectedly small")
    detailed = load_workbook(BytesIO(response.content), read_only=True)
    if "处理汇总" not in detailed.sheetnames or "原始数据" not in detailed.sheetnames:
        raise AssertionError(f"Detailed export has unexpected sheets: {detailed.sheetnames}")

    with sample_path.open("rb") as file_obj:
        marked_response = httpx.post(
            f"{APP_URL}/api/analyze-combined",
            files={"workbook": (sample_path.name, file_obj, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            data={"output_mode": "marked", "tolerance": "100"},
            timeout=30,
        )
    marked_response.raise_for_status()
    marked = load_workbook(BytesIO(marked_response.content), read_only=False)
    if marked.sheetnames != ["Sheet1"] or marked.active["A2"].fill.fgColor.rgb != "FFD2EEDA":
        raise AssertionError("Marked export did not color the original sheet")


def assert_ilog_local_services() -> None:
    # Never log in to iLog or submit business operations in CI.
    for endpoint in ('/api/ilog/runtime','/api/ilog/retention','/api/ilog/tasks?limit=20&offset=0'):
        response=httpx.get(APP_URL+endpoint,timeout=10)
        response.raise_for_status()
    headers={'X-Finance-Local':'1'}
    created=httpx.post(APP_URL+'/api/ilog/accounts',json={
        'name':'Smoke test','username':'smoke-test-only','password':'not-a-real-password'},headers=headers,timeout=10)
    created.raise_for_status()
    account=created.json()
    if not account.get('hasPassword') or 'not-a-real-password' in created.text:
        raise AssertionError('Account storage or secret redaction failed')
    response=httpx.delete(APP_URL+'/api/ilog/accounts/'+account['id'],headers=headers,timeout=10)
    response.raise_for_status()


def poll_json(path, deadline):
    """Cold native OCR startup may delay a request; keep a bounded overall wait."""
    last_error=None
    while time.monotonic()<deadline:
        try:
            response=httpx.get(APP_URL+path,headers={'X-Finance-Local':'1'},
                               timeout=max(.1,min(10,deadline-time.monotonic())))
            response.raise_for_status()
            return response.json()
        except httpx.TimeoutException as error:
            last_error=error
            time.sleep(.2)
    raise AssertionError('Frozen API did not respond before deadline: '+path) from last_error


def assert_mail_local_services() -> dict:
    from PIL import Image, ImageDraw, ImageFont
    # Check bundling only; CI must not connect to Outlook or any real mailbox.
    headers={'X-Finance-Local':'1'}
    for endpoint in ('/api/mail/status','/api/mail/messages'):
        response=httpx.get(APP_URL+endpoint,headers=headers,timeout=10)
        response.raise_for_status()
    if not isinstance(response.json(),list):
        raise AssertionError('Mail record storage is not available')
    for endpoint in ('/api/mail/online/config','/api/mail/online/history'):
        response=httpx.get(APP_URL+endpoint,headers=headers,timeout=10)
        response.raise_for_status()
        if not isinstance(response.json(),dict):raise AssertionError('Online ledger local storage is not available')
    image=Image.new('RGB',(1400,320),'white')
    draw=ImageDraw.Draw(image)
    font=ImageFont.load_default(size=40)
    draw.text((30,30),'Amount: IDR 1,400,000.00',font=font,fill='black')
    draw.text((30,110),'Value Date: 25-09-2026',font=font,fill='black')
    png=BytesIO(); image.save(png,format='PNG')
    pdf=BytesIO(); image.save(pdf,format='PDF')
    mail=EmailMessage();mail['Subject']='Synthetic offline OCR smoke test';mail.set_content('Test only')
    mail.add_attachment(png.getvalue(),maintype='image',subtype='png',filename='test.png')
    mail.add_attachment(pdf.getvalue(),maintype='application',subtype='pdf',filename='test.pdf')
    response=httpx.post(APP_URL+'/api/mail/import',headers=headers,files={'mail':('smoke.eml',mail.as_bytes())},timeout=30)
    response.raise_for_status()
    identity=response.json()['id']
    deadline=time.monotonic()+90
    while time.monotonic()<deadline:
        assets=poll_json('/api/mail/messages/'+identity,deadline)['attachment_details']
        if len(assets)==2 and all(a['status'] not in ('pending','running') for a in assets): break
        time.sleep(.2)
    else: raise AssertionError('Frozen attachment recognition timed out')
    for asset in assets:
        if asset['status']!='done': raise AssertionError(f"Frozen OCR/PDF engine failed: {asset.get('error')}")
        amounts=asset['result']['fields']['amount']
        if not any(f['value']=='1400000.00' for f in amounts):
            raise AssertionError('Frozen OCR failed to read the synthetic amount')

    before=wait_for_mail_analysis()
    duplicate=EmailMessage();duplicate['Subject']='Forwarded synthetic evidence';duplicate.set_content('Independent mail, same attachments')
    duplicate.add_attachment(png.getvalue(),maintype='image',subtype='png',filename='forwarded.png')
    duplicate.add_attachment(pdf.getvalue(),maintype='application',subtype='pdf',filename='forwarded.pdf')
    response=httpx.post(APP_URL+'/api/mail/import',headers=headers,files={'mail':('copy.eml',duplicate.as_bytes())},timeout=30)
    response.raise_for_status()
    after=wait_for_mail_analysis()
    previous=before['performance']['counters'];current=after['performance']['counters']
    if current['cache_hits']-previous['cache_hits']!=2 or current['analysis_runs']!=previous['analysis_runs']:
        raise AssertionError('Frozen duplicate attachments were not reused without new analysis')
    if current['ocr_pages']<2 or after['performance']['seconds']['analysis_seconds']<=0:
        raise AssertionError('Frozen performance recorder missed real OCR work')
    original=httpx.get(APP_URL+'/api/mail/messages/'+identity,headers=headers,timeout=10)
    original.raise_for_status()
    return dict(mail=original.json(),performance=after['performance'])


def wait_for_mail_analysis():
    deadline=time.monotonic()+90
    while time.monotonic()<deadline:
        status=poll_json('/api/mail/status',deadline)
        if status.get('analysis_error'):
            raise AssertionError('Frozen analysis failed: '+status['analysis_error'])
        if not status['analysis_busy']:
            if not status.get('performance'):
                raise AssertionError('Frozen performance recorder is unavailable')
            return status
        time.sleep(.2)
    raise AssertionError('Frozen mail queue did not drain')


def assert_mail_restart_preserved(expected):
    status=wait_for_mail_analysis()
    response=httpx.get(APP_URL+'/api/mail/messages/'+expected['id'],headers={'X-Finance-Local':'1'},timeout=10)
    response.raise_for_status()
    if response.json()!=expected:
        raise AssertionError('Restart changed stored mail, recognition results or payment evidence')
    if status['performance']['counters']['analysis_runs'] or status['analysis_pending']:
        raise AssertionError('Restart unnecessarily analyzed already completed attachments')


def start_test_process(exe, env):
    process=subprocess.Popen([str(exe)],env=env)
    if sys.platform=='win32':
        # A Windows job retains descendants even if the one-file bootloader exits.
        import ctypes
        from ctypes import wintypes
        kernel=ctypes.WinDLL('kernel32',use_last_error=True)
        kernel.CreateJobObjectW.argtypes=[ctypes.c_void_p,wintypes.LPCWSTR]
        kernel.CreateJobObjectW.restype=wintypes.HANDLE
        kernel.AssignProcessToJobObject.argtypes=[wintypes.HANDLE,wintypes.HANDLE]
        kernel.AssignProcessToJobObject.restype=wintypes.BOOL
        kernel.TerminateJobObject.argtypes=[wintypes.HANDLE,wintypes.UINT]
        kernel.TerminateJobObject.restype=wintypes.BOOL
        kernel.CloseHandle.argtypes=[wintypes.HANDLE]
        kernel.CloseHandle.restype=wintypes.BOOL
        job=kernel.CreateJobObjectW(None,None)
        if not job or not kernel.AssignProcessToJobObject(job,int(process._handle)):
            error=ctypes.WinError(ctypes.get_last_error())
            if job:kernel.CloseHandle(job)
            stop_process_tree(process)
            raise error
        process._smoke_job=(kernel,job)
    return process


def stop_process_tree(process):
    job=getattr(process,'_smoke_job',None)
    if job:
        kernel,handle=job
        try:
            if not kernel.TerminateJobObject(handle,1):
                import ctypes
                raise ctypes.WinError(ctypes.get_last_error())
            process.wait(timeout=10)
        finally:
            kernel.CloseHandle(handle)
            process._smoke_job=None
        return
    if process.poll() is not None:return
    if sys.platform=='win32':
        # PyInstaller one-file builds have a bootloader child; terminate the
        # complete test process tree before starting the same EXE again.
        result=subprocess.run(['taskkill','/PID',str(process.pid),'/T','/F'],capture_output=True,timeout=20)
        if result.returncode and process.poll() is None:
            raise AssertionError('Could not stop the frozen test process tree')
    else:
        process.terminate()
    try:process.wait(timeout=10)
    except subprocess.TimeoutExpired:
        process.kill();process.wait(timeout=10)


def assert_excel_isolation() -> None:
    for path in ('/api/ilog/accounts','/api/ilog/tasks','/api/mail/status','/api/mail/messages'):
        for method in ('GET','POST'):
            response=httpx.request(method,APP_URL+path,json={} if method=='POST' else None,timeout=10)
            if response.status_code!=404:raise AssertionError('Excel edition exposes forbidden API: '+path)


def main() -> int:
    excel_only='--excel-only' in sys.argv
    exe_path = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("dist") / "ilog解锁工具.exe"
    if not exe_path.exists():
        raise FileNotFoundError(exe_path)

    try:
        httpx.get(f"{APP_URL}/api/health", timeout=2)
    except httpx.RequestError:
        pass
    else:
        raise RuntimeError("Port 8765 is already serving an app; cannot verify the new EXE")

    # Verify imports and the bundled Node bridge inside the EXE, rather than
    # accidentally testing the CI runner's installed Python packages.
    diagnostics=Path('build')/'smoke-diagnostics'
    diagnostics.mkdir(parents=True,exist_ok=True)
    startup_log=(diagnostics/'startup.log').resolve()
    startup_log.write_text('',encoding='utf-8')
    process_env={**os.environ,'FINANCE_RECONCILE_STARTUP_LOG':str(startup_log)}
    with tempfile.TemporaryDirectory(prefix='finance-bundle-probe-') as probe_dir:
        report=Path(probe_dir)/'probe.json'
        checked=subprocess.run([str(exe_path),'--verify-bundle',str(report)],timeout=90,env=process_env)
        if not report.exists():raise AssertionError('Frozen runtime probe did not produce a report')
        result=json.loads(report.read_text(encoding='utf-8'))
        (diagnostics/'probe.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
        if checked.returncode or not result.get('ok'):
            raise AssertionError(f'Frozen runtime components failed: {result}\n{startup_log.read_text(encoding="utf-8",errors="replace")[-24000:]}')

    sandbox = tempfile.TemporaryDirectory(prefix='finance-exe-smoke-')
    app_env={**process_env,'FINANCE_RECONCILE_DATA_DIR':sandbox.name}
    process = subprocess.Popen([str(exe_path)], env=app_env)
    try:
        wait_for_health(process=process,startup_log=startup_log)
        assert_frontend_served()
        if excel_only:
            assert_excel_isolation()
        else:
            assert_ilog_local_services()
            sample=assert_mail_local_services()
            (diagnostics/'mail-performance.json').write_text(json.dumps(sample['performance'],indent=2),encoding='utf-8')
            key=Path(sandbox.name)/'mail'/'mail.key'
            key_before=key.read_bytes()
            stop_process_tree(process)
            process=subprocess.Popen([str(exe_path)],env=app_env)
            wait_for_health(process=process,startup_log=startup_log)
            assert_mail_restart_preserved(sample['mail'])
            if key.read_bytes()!=key_before:
                raise AssertionError('Restart replaced the original mail encryption key')
            (diagnostics/'mail-restart.json').write_text(json.dumps(dict(ok=True,completed_results_preserved=True,key_preserved=True)),encoding='utf-8')
        with tempfile.TemporaryDirectory() as temp_dir:
            sample_path = Path(temp_dir) / "sample.xlsx"
            make_sample_workbook(sample_path)
            assert_excel_flow(sample_path)
    finally:
        stop_process_tree(process)
        sandbox.cleanup()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
