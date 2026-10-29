import pytest
from app.mail_matching import rank_receipts
from app.customer_mapping import MappingIndex, parse_contacts


def entry(**changes):
    result=dict(customer='FAF',bank_name='Nani Desman',emails=['dian@faf.test'],row=3,sheet='Sheet1')
    result.update(changes);return result


def payment(**changes):
    result=dict(id='p',mail_id='m',attachment_id='a',payer=None,date='2026-09-20',amount='100.00',currency='IDR',sender='Dian <dian@faf.test>',warnings=[],evidence={},remark='original',search_info='',business_ids=[])
    result.update(changes);return result


def receipt(**changes):
    result=dict(id='r',payer='Nani Desman',date='2026-09-20',amount='100.00',currency='IDR')
    result.update(changes);return result


def ranked(pay=None,rc=None,entries=None):
    return rank_receipts([rc or receipt()],[pay or payment()],mapping=MappingIndex(entries or [entry()]))[0]['candidates']


def test_mapping_improves_missing_payer_and_preserves_source():
    source=payment();plain=rank_receipts([receipt()],[source])[0]['candidates'][0]
    mapped=ranked(source)[0]
    assert mapped['score']>plain['score'] and mapped['payer'] is None
    assert mapped['breakdown']['mapping']==24
    assert mapped['mapping']['rows'][0]['bank_name']=='Nani Desman'
    assert '客户台账' in mapped['remark'] and source['remark']=='original'
    assert source['evidence']=={}


def test_explicit_company_alias_and_stronger_direct_name():
    alias=ranked(payment(payer='FAF'))[0]
    assert alias['breakdown']['mapping']==22
    direct=ranked(payment(payer='Nani Desman'))[0]
    assert direct['breakdown']['payer']==30 and direct['breakdown'].get('mapping',0)==0
    assert not direct.get('mapping')


def test_unrelated_payer_or_sender_not_overridden():
    assert not ranked(payment(payer='Other Person'))[0].get('mapping')
    assert not ranked(payment(sender='Dian <other@faf.test>'))[0].get('mapping')


@pytest.mark.parametrize('changes',[dict(currency='USD'),dict(currency=''),dict(amount='99.00'),dict(date='2025-01-01')])
def test_mapping_cannot_override_money_currency_date_conflicts(changes):
    assert all(not p.get('mapping') for p in ranked(payment(**changes)))


def test_shared_mailbox_gets_lower_score_and_ambiguity_warning():
    rows=[entry(),entry(customer='Other Customer',bank_name='Other Payer',row=4)]
    mapped=ranked(entries=rows)[0]
    assert mapped['breakdown']['mapping']==10 and mapped['mapping']['ambiguous']
    assert any('共享' in w for w in mapped['warnings'])


def test_contact_csv_display_names_and_incomplete_rows():
    raw=('No,Customer Name,Customer Bank Name Reff,Email\n1,FAF,Nani Desman,Dian <dian@faf.test>\n2,Other,,other@test.com\n').encode()
    result=parse_contacts(raw,'contacts.csv')
    assert result['entries'][0]['emails']==['dian@faf.test']
    assert result['skipped']==1
    with pytest.raises(ValueError):parse_contacts(b'wrong,headers\na,b','bad.csv')


def test_grouped_duplicate_sources_use_all_sender_relationships():
    payments=[payment(id='first',mail_id='first',sender='unrelated@test.com',attachment_sha256='same'),
              payment(id='second',mail_id='second',attachment_sha256='same')]
    ranked_rows=rank_receipts([receipt()],payments,mapping=MappingIndex([entry()]))
    assert len(ranked_rows[0]['candidates'])==1
    assert ranked_rows[0]['candidates'][0]['breakdown']['mapping']==24


def csv_contacts():
    return b'Customer Name,Customer Bank Name Reff,Email\nFAF,Nani Desman,dian@faf.test\n'


def workspace(tmp_path):
    from app.outlook_probe import MailStore
    from app.mail_workspace import MailWorkspace
    store=MailStore(tmp_path/'mail.sqlite3');ws=MailWorkspace(store)
    store.save([dict(id='m',subject='Original')])
    store.set_payments('m',[payment()])
    ws.import_ledger(b'DATE,PAYER,CREDIT,CURRENCY,REF NO\n2026-09-20,Nani Desman,100.00,IDR,001\n','bank.csv')
    return ws


def test_mapping_import_cache_confirmation_export_and_disable(tmp_path):
    from app.mail_workspace import signature
    from openpyxl import load_workbook
    from io import BytesIO
    ws=workspace(tmp_path)
    original=ws.matching(scope='all')['rows'][0]['candidates'][0]
    metadata=ws.import_customer_mapping(csv_contacts(),'customers.csv')
    assert metadata['entries_count']==1 and 'entries' not in metadata
    assert b'Nani Desman' not in ws.store.path.read_bytes()
    changed=ws.matching(scope='all');row=changed['rows'][0];candidate=row['candidates'][0]
    assert candidate['score']>original['score']
    assert candidate['evidence_signature']==signature(ws._candidate(row['receipt'],candidate['payment_id'],3))
    with pytest.raises(ValueError):ws.review(changed['ledger']['id'],row['receipt']['id'],'p','confirmed',3,original['evidence_signature'])
    ws.review(changed['ledger']['id'],row['receipt']['id'],'p','confirmed',3,candidate['evidence_signature'])
    assert ws.matching(scope='all')['summary']['confirmed']==1
    raw,name=ws.export(scope='all')
    book=load_workbook(BytesIO(raw));assert any('客户台账' in str(cell.value) for row in book.active for cell in row);book.close()
    ws.enable_customer_mapping(False)
    plain=ws.matching(scope='all')
    assert plain['summary']['confirmed']==0
    assert not plain['rows'][0]['candidates'][0].get('mapping')
    assert ws.store.payment_rows()[0]['payer'] is None
    assert ws.store.get('m')['subject']=='Original'


def test_invalid_import_keeps_previous_mapping_and_reopen(tmp_path):
    from app.mail_workspace import MailWorkspace
    ws=workspace(tmp_path);old=ws.import_customer_mapping(csv_contacts(),'customers.csv')
    with pytest.raises(ValueError):ws.import_customer_mapping(b'wrong,columns\n1,2','bad.csv')
    assert ws.customer_mapping()==old
    assert MailWorkspace(ws.store).customer_mapping()==old


def test_xlsx_header_offset_and_selected_sheet(tmp_path):
    from io import BytesIO
    from openpyxl import Workbook
    book=Workbook();book.active.title='Notes';sheet=book.create_sheet('Customers')
    sheet.append(['Customer contacts']);sheet.append(['Customer Name','Customer Bank Name Reff','Email'])
    sheet.append(['FAF','Nani Desman','Dian <dian@faf.test>']);out=BytesIO();book.save(out)
    result=parse_contacts(out.getvalue(),'contacts.xlsx','Customers')
    assert result['entries'][0]['row']==3
    assert result['entries'][0]['emails']==['dian@faf.test']
    with pytest.raises(ValueError):parse_contacts(out.getvalue(),'contacts.xlsx','missing')


def test_api_import_toggle_and_invalid_file(tmp_path,monkeypatch):
    from fastapi.testclient import TestClient
    from app import mail_tasks
    from app.main import app
    ws=workspace(tmp_path)
    manager=mail_tasks.MailManager(ws.store,reader=lambda *x:dict(rows=[],errors=[],partial=False))
    monkeypatch.setattr(mail_tasks,'_manager',manager)
    client=TestClient(app);headers={'X-Finance-Local':'1'}
    assert client.get('/api/mail/customer-mapping').status_code==403
    response=client.post('/api/mail/customer-mapping/import',headers=headers,files={'workbook':('customers.csv',csv_contacts())})
    assert response.status_code==200 and response.json()['enabled']
    assert client.put('/api/mail/customer-mapping',headers=headers,json={'enabled':False}).json()['enabled'] is False
    assert client.post('/api/mail/customer-mapping/import',headers=headers,files={'workbook':('bad.csv',b'wrong,header')}).status_code==400
    assert client.get('/api/mail/customer-mapping',headers=headers).json()['entries_count']==1


def test_direct_confirmation_survives_unrelated_mapping_import(tmp_path):
    ws=workspace(tmp_path);ws.store.set_payments('m',[payment(payer='Nani Desman')])
    initial=ws.matching(scope='all');row=initial['rows'][0];candidate=row['candidates'][0]
    ws.review(initial['ledger']['id'],row['receipt']['id'],'p','confirmed',3,candidate['evidence_signature'])
    ws.import_customer_mapping(csv_contacts(),'customers.csv')
    assert ws.matching(scope='all')['summary']['confirmed']==1


def test_mapping_reorders_candidates_but_direct_match_stays_first():
    payments=[payment(id='other',sender='other@faf.test'),payment(id='mapped'),payment(id='direct',payer='Nani Desman')]
    result=rank_receipts([receipt()],payments,mapping=MappingIndex([entry()]))
    assert [p['id'] for p in result[0]['candidates']]==['direct','mapped','other']


def test_same_company_shared_by_multiple_bank_payers_is_ambiguous():
    result=ranked(entries=[entry(),entry(bank_name='Another Person',row=4)])[0]
    assert result['mapping']['ambiguous'] and result['breakdown']['mapping']==10


def test_reverse_lookup_and_confirmation_agree_for_grouped_sources(tmp_path):
    from app.mail_insights import related_receipts
    ws=workspace(tmp_path);ws.import_customer_mapping(csv_contacts(),'customers.csv')
    ws.store.set_payments('m',[payment(sender='other@test.com',attachment_sha256='same')])
    ws.store.save([dict(id='copy')]);ws.store.set_payments('copy',[payment(id='copy-payment',mail_id='copy',attachment_sha256='same')])
    forward=ws.matching(scope='all')['rows'][0]['candidates'][0]
    reverse=related_receipts(ws,'copy')['rows'][0]['candidates'][0]
    assert forward['payment_id']==reverse['payment_id']
    assert forward['score']==reverse['score']
    assert forward['evidence_signature']==reverse['evidence_signature']


def test_equal_unknown_currency_cannot_activate_mapping():
    for code in ('UNK','UNKNOWN','XXX'):
        assert not ranked(payment(currency=code),receipt(currency=code))[0].get('mapping')


def test_existing_confirmation_of_noncanonical_copy_keeps_original_evidence(tmp_path):
    from app.mail_workspace import signature
    ws=workspace(tmp_path)
    ws.store.set_payments('m',[payment(payer='Nani Desman',attachment_sha256='same')])
    copy=payment(id='copy-payment',mail_id='copy',payer='Nani Desman',attachment_sha256='same',subject='Archived copy')
    ws.store.save([dict(id='copy')]);ws.store.set_payments('copy',[copy])
    before=ws.matching(scope='all');row=before['rows'][0]
    legacy=rank_receipts([row['receipt']],[copy])[0]['candidates'][0]
    with ws.store.db() as db:
        db.execute('INSERT INTO mail_review VALUES(?,?)',(row['receipt']['id'],ws.store.encode(dict(state='confirmed',payment_id='copy-payment',signature=signature(legacy)))))
    after=ws.matching(scope='all')
    assert after['summary']['confirmed']==1
    selected=next(c for c in after['rows'][0]['candidates'] if c['payment_id']=='copy-payment')
    assert selected['evidence_signature']==signature(legacy)
