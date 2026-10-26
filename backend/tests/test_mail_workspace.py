from io import BytesIO

import openpyxl
import pytest

from app.outlook_probe import MailStore


def ledger_bytes():
    workbook=openpyxl.Workbook();sheet=workbook.active;sheet.title='Bank Transactions'
    sheet.append(['NO','EFF DATE','DESCRIPTION','CREDIT','BUSINESS TYPE','REF NO','CLOSING STATUS'])
    sheet.append([1,'25 Sep 2026','BFS FR BANK UN TEST PAYER',1400000,'FUEL','001234567890','UNIDENTIFIED'])
    sheet.append([2,'25 Sep 2026','BFS FR BANK UN TEST PAYER',1400000,'FUEL','closed-ref','CLOSED'])
    out=BytesIO();workbook.save(out);return out.getvalue()


def payment():
    return dict(id='payment',payment_id='payment',mail_id='mail',attachment_id='attachment',page=1,
                payer='TEST PAYER',date='2026-09-25',amount='1400000.00',currency='IDR',
                subject='Payment TESTBL123456',sender='sender@example.invalid',received='2026-09-26T12:00:00+08:00',
                keyword='TESTBL123456',business_ids=['TESTBL123456'],remark='=Not a formula',
                search_info='主题：Payment TESTBL123456\n发件人：sender@example.invalid',evidence={},warnings=[],duplicate_key='evidence')


def test_indexed_mail_matches_review_and_preserves_bank_state(tmp_path):
    from app.mail_workspace import MailWorkspace
    store=MailStore(tmp_path/'mail.sqlite3');workspace=MailWorkspace(store)
    metadata=workspace.import_ledger(ledger_bytes(),'bank.xlsx')
    assert metadata['entries_count']==2
    store.set_payments('mail',[payment()])
    result=workspace.matching()
    assert len(result['rows'])==1
    row=result['rows'][0]
    assert row['receipt']['ref']=='001234567890'
    assert row['candidates'][0]['score']==100
    workspace.review(metadata['id'],row['receipt']['id'],'payment','confirmed',3)
    assert workspace.matching()['rows'][0]['review']['state']=='confirmed'
    raw=workspace.export()[0]
    exported=openpyxl.load_workbook(BytesIO(raw));sheet=exported['Bank Transactions']
    assert sheet['G2'].value=='UNIDENTIFIED' and sheet['G3'].value=='CLOSED'
    assert sheet['H2'].data_type=='s'
    assert '=Not a formula' in sheet['H2'].value
    assert sheet['I2'].value==100
    assert 'sender@example.invalid' in sheet['J2'].value
    assert sheet['H3'].value is None


def test_review_rejects_replaced_ledger_and_unrelated_candidate(tmp_path):
    from app.mail_workspace import MailWorkspace
    workspace=MailWorkspace(MailStore(tmp_path/'mail.sqlite3'))
    workspace.import_ledger(ledger_bytes(),'bank.xlsx')
    receipt=workspace.matching()['rows'][0]['receipt']['id']
    with pytest.raises(ValueError): workspace.review('wrong-ledger',receipt,None,'rejected',3)
    with pytest.raises(ValueError): workspace.review(workspace.ledger()['id'],receipt,'missing','confirmed',3)


def test_changed_payment_requires_review_again(tmp_path):
    from app.mail_workspace import MailWorkspace
    store=MailStore(tmp_path/'mail.sqlite3');workspace=MailWorkspace(store)
    meta=workspace.import_ledger(ledger_bytes(),'bank.xlsx');store.set_payments('mail',[payment()])
    receipt=workspace.matching()['rows'][0]['receipt']['id']
    workspace.review(meta['id'],receipt,'payment','confirmed',3)
    changed={**payment(),'payer':'OTHER PAYER','remark':'different evidence'}
    store.set_payments('mail',[changed])
    row=workspace.matching()['rows'][0]
    assert row['review']['state']=='pending'
    assert any('更新' in warning for warning in row['warnings'])


def test_matching_routes_are_local_and_ledger_is_persistent(tmp_path,monkeypatch):
    from fastapi.testclient import TestClient
    from app.main import app
    from app import mail_tasks
    store=MailStore(tmp_path/'mail.sqlite3')
    monkeypatch.setattr(mail_tasks,'_manager',mail_tasks.MailManager(store))
    client=TestClient(app);headers={'X-Finance-Local':'1','Origin':'http://127.0.0.1:5173'}
    assert client.get('/api/mail/ledger').status_code==403
    response=client.post('/api/mail/ledger/import',headers=headers,files={'workbook':('bank.xlsx',ledger_bytes())},data={'currency':'IDR'})
    assert response.status_code==200
    assert client.get('/api/mail/matching',headers=headers).json()['total_rows']==1
    assert client.get('/api/mail/matching?scope=all',headers=headers).json()['total_rows']==2
    assert client.get('/api/mail/matching/export',headers=headers).headers['content-type'].startswith('application/vnd.openxmlformats')
    from app.mail_workspace import MailWorkspace
    assert MailWorkspace(MailStore(tmp_path/'mail.sqlite3')).ledger()['id']==response.json()['id']


def test_semicolon_csv_export_keeps_zero_padded_reference(tmp_path):
    from app.mail_workspace import MailWorkspace
    workspace=MailWorkspace(MailStore(tmp_path/'mail.sqlite3'))
    raw=b'DATE;AMOUNT;REF NO;CLOSING STATUS\n2026-09-25;1400000;000123;UNIDENTIFIED\n'
    workspace.import_ledger(raw,'bank.csv')
    data,name=workspace.export()
    sheet=openpyxl.load_workbook(BytesIO(data)).active
    assert sheet['C2'].value=='000123'
    assert sheet['D2'].value=='UNIDENTIFIED'


def test_search_filters_full_ledger_before_pagination_and_score_order(tmp_path):
    from app.mail_workspace import MailWorkspace
    store=MailStore(tmp_path/'mail.sqlite3');workspace=MailWorkspace(store)
    workspace.import_ledger(ledger_bytes(),'bank.xlsx');store.set_payments('mail',[payment()])
    result=workspace.matching(scope='all',page_size=1,keyword='closed-ref',sort_by='score')
    assert result['total_rows']==1 and result['rows'][0]['receipt']['ref']=='closed-ref'
    assert workspace.matching(scope='all',keyword='nonexistent')['total_rows']==0
    assert workspace.matching(scope='all',keyword='test payer')['total_rows']==2


def test_score_sort_puts_stronger_candidate_before_original_row_order(tmp_path):
    from app.mail_workspace import MailWorkspace
    workbook=openpyxl.load_workbook(BytesIO(ledger_bytes()));workbook.active['C2']='BFS FR BANK UN OTHER PERSON'
    raw=BytesIO();workbook.save(raw)
    store=MailStore(tmp_path/'mail.sqlite3');workspace=MailWorkspace(store)
    workspace.import_ledger(raw.getvalue(),'bank.xlsx');store.set_payments('mail',[payment()])
    assert workspace.matching(scope='all')['rows'][0]['receipt']['ref']=='001234567890'
    assert workspace.matching(scope='all',sort_by='score')['rows'][0]['receipt']['ref']=='closed-ref'
