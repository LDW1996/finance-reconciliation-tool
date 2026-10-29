from fastapi.testclient import TestClient
from app.mail_main import app
from app.mail_english import response

def test_mail_route_isolation():
    client=TestClient(app)
    assert client.get('/api/health').json()['edition']=='mail_en'
    for path in ['/api/ilog/accounts','/api/reconcile','/api/excel/status']:
        assert client.get(path).status_code==404
    assert client.get('/api/mail/status').status_code==403

def test_preserve_business_text():
    value={'subject':'完成','body':'付款人缺失，请查看凭证。','payer':'失败','warnings':['付款人缺失，请查看凭证。']}
    translated=response(value)
    assert translated['subject']==value['subject']
    assert translated['body']==value['body']
    assert translated['payer']==value['payer']
    assert translated['warnings']==['Missing payer']
    assert value['warnings']==['付款人缺失，请查看凭证。']

def test_export_preserves_original_chinese_cells():
    from io import BytesIO
    from openpyxl import Workbook,load_workbook
    from app.mail_english import english_export
    book=Workbook();sheet=book.active
    sheet.append(['银行原文','邮件关联备注','匹配分数','邮件检索信息'])
    sheet.append(['邮件主题：中文客户','未匹配到候选邮件',None,'邮件主题：中文客户'])
    source=BytesIO();book.save(source)
    result=load_workbook(BytesIO(english_export(source.getvalue()))).active
    assert result['A2'].value=='邮件主题：中文客户'
    assert result['B1'].value=='Mail association notes'
    assert result['D2'].value=='Mail subject:中文客户'

def test_same_mail_storage_directory(monkeypatch,tmp_path):
    from app.outlook_probe import data_path
    monkeypatch.setenv('FINANCE_RECONCILE_DATA_DIR',str(tmp_path))
    assert data_path()==tmp_path/'mail'/'mail.sqlite3'

def test_path_aware_labels_and_preview_diagnostics():
    original={'folders':[{'label':'共享邮箱'}], 'mapping':{'label':'客户台账映射（共享关系，需核对）'},
              'preview':{'flags':['共享邮箱','一对多关系'],'reasons':['客户名称缺失'],'notice':'客户名称缺失'},
              'customer_mapping':{'label':'客户台账映射'},'customer':'共享邮箱'}
    result=response(original)
    assert result['folders'][0]['label']=='共享邮箱'
    assert result['customer']=='共享邮箱'
    assert result['mapping']['label']=='Customer ledger mapping (shared relationship; review required)'
    assert result['preview']['flags']==['Shared email','One-to-many relationship']
    assert result['preview']['reasons']==['Customer name is missing']
    assert original['preview']['flags']==['共享邮箱','一对多关系']

def test_structured_payment_and_mapping_preserve_business_names():
    from app.mail_english import generated_lines
    raw='付款：日期未识别；币种未识别 金额未识别；中文付款人\n邮件主题：中文客户付款\n客户台账关联（辅助依据）：中文付款人 → 中文客户；邮箱：未登记；中文工作表 第12行'
    result=generated_lines(raw)
    for generated in ['付款：','日期未识别','币种未识别','金额未识别','邮件主题：','；邮箱：','第12行']:
        assert generated not in result
    for business in ['中文付款人','中文客户付款','中文客户','中文工作表']:
        assert business in result
    assert 'Email: Not registered' in result
    assert 'row 12' in result

def test_failure_diagnostic_meaningful_english():
    value={'analysis_diagnostic':{'kind':'disk_full','phase':'检查存储空间','reason':'存储空间不足',
          'action':'释放磁盘空间后手动重试','message':'检查存储空间失败','retryable':False}}
    result=response(value)['analysis_diagnostic']
    assert result['phase']=='Checking storage space'
    assert result['message']=='Checking storage space: Insufficient disk space. Free disk space and retry manually'


def test_english_preview_can_search_translated_flags_without_changing_business_text(tmp_path,monkeypatch):
    from app import mail_tasks
    from app.outlook_probe import MailStore
    from app.mail_workspace import MailWorkspace
    from types import SimpleNamespace
    store=MailStore(tmp_path/'mail.sqlite3');workspace=MailWorkspace(store)
    workspace.import_customer_mapping('Customer Name,Customer Bank Name Reff,Email\n上海公司,Person A,same@test.com\n深圳客户,Person B,same@test.com\n缺少姓名,,bad@test.com\n'.encode(),'contacts.csv')
    monkeypatch.setattr(mail_tasks,'_manager',SimpleNamespace(store=store))
    client=TestClient(app);headers={'X-Finance-Local':'1'}
    result=client.get('/api/mail/customer-mapping/preview',params={'keyword':'Shared email'},headers=headers).json()
    assert result['total']==2 and result['rows'][0]['customer']=='上海公司'
    assert result['rows'][0]['flags']==['Shared email']
    result=client.get('/api/mail/customer-mapping/preview',params={'status':'excluded','keyword':'Bank payer is missing'},headers=headers).json()
    assert result['total']==1 and result['rows'][0]['customer']=='缺少姓名'
    assert result['rows'][0]['reasons']==['Bank payer is missing']


def test_english_writeback_generates_english_cells_and_preserves_bank_values(tmp_path):
    from test_mail_online import setup_online
    from app.mail_online import OUTPUT_LANGUAGE
    online,workspace,client=setup_online(tmp_path)
    original=[list(row) for row in client.values]
    token=OUTPUT_LANGUAGE.set('en')
    try:result=online.writeback(workspace.ledger()['id'],3)
    finally:OUTPUT_LANGUAGE.reset(token)
    assert result['written']==1
    assert client.values[0][-3:]==['Mail association notes','Match score','Mail search information']
    assert client.values[1][:len(original[1])]==original[1]
    assert client.values[1][-3].startswith('Confirmed association:')
    assert OUTPUT_LANGUAGE.get()=='zh-CN'


def test_forbidden_api_methods_are_not_registered():
    client=TestClient(app)
    for method in ('get','post','put','delete'):
        assert getattr(client,method)('/api/ilog/accounts').status_code==404
        assert getattr(client,method)('/api/analyze-combined').status_code==404


def test_packaging_artifact_names_and_guide_agree(monkeypatch):
    import ast
    from pathlib import Path
    # Exercise Windows' non-UTF-8 default even on UTF-8 development machines.
    original_read_text=Path.read_text
    def windows_read_text(path,encoding=None,errors=None):
        return original_read_text(path,encoding=encoding or 'cp1252',errors=errors)
    monkeypatch.setattr(Path,'read_text',windows_read_text)
    root=Path(__file__).resolve().parents[2]
    tree=ast.parse((root/'packaging/mail-reconcile.spec').read_text(encoding='utf-8'))
    names=[kw.value.value for node in ast.walk(tree) if isinstance(node,ast.Call) and isinstance(node.func,ast.Name) and node.func.id=='EXE' for kw in node.keywords if kw.arg=='name']
    assert names==['邮件核销工具（英文版）']
    assert names[0] in (root/'.github/workflows/build-windows-exe.yml').read_text(encoding='utf-8')
    guide=(root/'packaging/Mail-Reconciliation-Guide.txt').read_text(encoding='utf-8')
    assert names[0]+'.exe' in guide
