from io import BytesIO
import sys
import subprocess
import os

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook


def test_excel_edition_exports_workbook_but_has_no_management_routes():
    from app.excel_main import app
    client=TestClient(app)
    assert client.get('/api/health').json()['status']=='ok'
    for path in ['/api/ilog/accounts','/api/ilog/tasks','/api/mail/status','/api/mail/messages']:
        assert client.get(path).status_code==404
        assert client.post(path,json={}).status_code==404
    book=Workbook();sheet=book.active
    sheet.append(['公司代码or伙伴公司','分配','原币金额'])
    sheet.append(['1703','A',100]);sheet.append(['2615','A',-100])
    raw=BytesIO();book.save(raw)
    for mode in ['detailed','marked']:
        response=client.post('/api/analyze-combined',files={'workbook':('sample.xlsx',raw.getvalue())},data={'output_mode':mode,'locale':'id-ID'})
        assert response.status_code==200
        assert load_workbook(BytesIO(response.content)).sheetnames


def test_excel_server_does_not_import_full_product_modules():
    code="import sys; import app.excel_main; assert not any(x in sys.modules for x in ['app.main','app.accounts','app.ilog_tasks','app.mail_tasks','playwright','onnxruntime'])"
    subprocess.run([sys.executable,'-c',code],check=True,env={**os.environ,'PYTHONPATH':'backend'})
