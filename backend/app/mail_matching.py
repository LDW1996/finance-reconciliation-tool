"""Pure receipt-centric matching. Scores are similarity points, not probabilities.

Missing scalar payment fields are None. Ambiguous pages are retained with all
scoreable fields cleared; evidence remains available for human inspection.
"""

import csv
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from difflib import SequenceMatcher
from hashlib import sha256
from io import BytesIO, StringIO
from itertools import islice
from pathlib import Path
import re
import unicodedata

from dateutil import parser as date_parser


MAX_BYTES = 25 * 1024 * 1024
MAX_ROWS = 20_000
MAPPING_CURRENCIES = frozenset(('IDR','USD','CNY','SGD','EUR','HKD','MYR','THB','JPY','AUD','GBP'))
_INDONESIAN_MONTHS={'januari':'Jan','februari':'Feb','maret':'Mar','mei':'May','juni':'Jun','juli':'Jul',
                    'agustus':'Aug','agu':'Aug','agt':'Aug','oktober':'Oct','okt':'Oct','desember':'Dec','des':'Dec'}


def _text(value):
    return '' if value is None else str(value).strip()


def _currency(value):
    value = _text(value).upper()
    return 'IDR' if value in ('RP', 'RP.') else value


def _amount(value):
    """Accept valid English/Indonesian grouping, never arbitrary punctuation."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (float, int, Decimal)):
        text = str(value)
    else:
        text = _text(value).replace('\u00a0', ' ')
        text = re.sub(r'^(?:IDR|RP\.?|USD|EUR|CNY)\s*|\s*(?:IDR|USD|EUR|CNY)$', '', text, flags=re.I)
        if re.fullmatch(r'[+-]?\d{1,3}(?: \d{3})+(?:[.,]\d{1,2})?', text):
            text = text.replace(' ', '')
        if not re.fullmatch(r'[+-]?\d+(?:\.\d{1,2})?', text):
            if re.fullmatch(r'[+-]?\d{1,3}(?:,\d{3})+(?:\.\d{1,2})?', text):
                text = text.replace(',', '')
            elif re.fullmatch(r'[+-]?\d{1,3}(?:\.\d{3})+(?:,\d{1,2})?', text):
                text = text.replace('.', '').replace(',', '.')
            elif re.fullmatch(r'[+-]?\d+,\d{1,2}', text):
                text = text.replace(',', '.')
            else:
                return None
    try:
        number = Decimal(text)
        if not number.is_finite() or number != number.quantize(Decimal('.01')):
            return None
        return format(number.quantize(Decimal('.01')), 'f')
    except InvalidOperation:
        return None


def _date(value):
    if isinstance(value, (datetime, date)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    text = _text(value)
    if not text or text.isdigit():
        return None
    text = re.sub(r'\bSept\b', 'Sep', text, flags=re.I)
    text = re.sub(r'\b(?:'+ '|'.join(_INDONESIAN_MONTHS) +r')\b',lambda match:_INDONESIAN_MONTHS[match[0].lower()],text,flags=re.I)
    text = re.sub(r'[年/月]', '-', text).replace('日', '')
    try:
        # Require all calendar components: dateutil must not supply today's date.
        first = date_parser.parse(text, dayfirst=not bool(re.match(r'^\d{4}[-/]', text)), default=datetime(2000, 1, 1))
        second = date_parser.parse(text, dayfirst=not bool(re.match(r'^\d{4}[-/]', text)), default=datetime(2004, 2, 2))
        return first.date().isoformat() if first.date() == second.date() else None
    except (ValueError, OverflowError, TypeError):
        return None


def _name(value):
    text = unicodedata.normalize('NFKC', _text(value)).upper()
    text = re.sub(r'[^\w]', '', text, flags=re.UNICODE).replace('_', '')
    # Corporate prefixes also disappear when OCR removes every space.
    text = re.sub(r'^(?:PT|CV|PERSEROANTERBATAS)', '', text)
    return text


_ALIASES = {
    'date': ('EFF DATE', 'EFFECTIVE DATE', 'VALUE DATE', 'TRANSACTION DATE', 'DATE', '日期', '交易日期', '入账日期', '到账日期'),
    'amount': ('CREDIT', 'CREDIT AMOUNT', 'AMOUNT', '金额', '收入', '收入金额', '贷方金额', '收款金额', '到账金额'),
    'payer': ('PAYER', 'PAYER NAME', 'REMITTER', '付款人', '付款单位', '付款方', '汇款人'),
    'ref': ('REF NO', 'REFERENCE', 'REFERENCE NO', 'REF', '流水号', '参考号', '交易流水号', '凭证号'),
    'description': ('DESCRIPTION', 'REMARK', 'REMARKS', '摘要', '描述', '备注'),
    'status': ('CLOSING STATUS', 'STATUS', '状态', '核销状态', '结算状态'),
    'currency': ('CURRENCY', 'CCY', '币种', '货币'),
}


def _header(value):
    return re.sub(r'[\W_]', '', _text(value)).upper()


def parse_ledger(raw: bytes, filename: str, sheet_name: str | None = None, currency='IDR') -> dict:
    """Import at most 20,000 data rows; row/header_row are one-based source rows.

    All statuses, including closed rows, are imported. Invalid date/amount rows
    are skipped with counts and source-row warnings. No currency conversion.
    """
    if len(raw) > MAX_BYTES:
        raise ValueError('Ledger exceeds 25 MB')
    extension = Path(filename).suffix.lower()
    book = None
    try:
        if extension in ('.xlsx', '.xlsm'):
            from openpyxl import load_workbook
            book = load_workbook(BytesIO(raw), read_only=True, data_only=True)
            sheets = list(book.sheetnames)
            selected = sheet_name or next((s for s in sheets if s.lower() == 'bank transactions'), sheets[0])
            if selected not in sheets:
                raise ValueError('Unknown ledger sheet')
            def cells():
                for row in book[selected].iter_rows():
                    values = []
                    for cell in row:
                        value = cell.value
                        if isinstance(value, (int, float)) and re.fullmatch(r'0+', cell.number_format or ''):
                            value = f'{int(value):0{len(cell.number_format)}d}' if value == int(value) else value
                        values.append(value)
                    yield values
            rows = cells()
        elif extension == '.xls':
            import pandas as pd
            book = pd.ExcelFile(BytesIO(raw), engine='xlrd')
            sheets = list(book.sheet_names)
            selected = sheet_name or next((s for s in sheets if s.lower() == 'bank transactions'), sheets[0])
            if selected not in sheets:
                raise ValueError('Unknown ledger sheet')
            frame = pd.read_excel(book, sheet_name=selected, header=None, dtype=object, keep_default_na=False, nrows=MAX_ROWS + 21)
            # xlrd retains numeric reference display masks separately from pandas.
            import xlrd
            legacy = xlrd.open_workbook(file_contents=raw, formatting_info=True, on_demand=True)
            try:
                source = legacy.sheet_by_name(selected)
                for r in range(min(len(frame), source.nrows)):
                    for c in range(min(len(frame.columns), source.ncols)):
                        cell = source.cell(r, c)
                        mask = legacy.format_map[legacy.xf_list[cell.xf_index].format_key].format_str
                        if cell.ctype == xlrd.XL_CELL_NUMBER and re.fullmatch(r'0+', mask) and cell.value == int(cell.value):
                            frame.iat[r, c] = f'{int(cell.value):0{len(mask)}d}'
            finally:
                legacy.release_resources()
            rows = frame.itertuples(index=False, name=None)
        elif extension == '.csv':
            try:
                decoded = raw.decode('utf-8-sig')
            except UnicodeDecodeError:
                decoded = raw.decode('gb18030')
            try:
                dialect = csv.Sniffer().sniff(decoded[:65536], delimiters=',;\t|')
            except csv.Error:
                dialect = csv.excel
            rows = csv.reader(StringIO(decoded), dialect)
            sheets, selected = [], None
        else:
            raise ValueError('Supported ledgers: xlsx, xls, csv')
        rows = iter(rows)
        aliases = {_header(alias): key for key, names in _ALIASES.items() for alias in names}
        columns = {}
        for header_row, values in enumerate(islice(rows, 20), 1):
            columns = {}
            for i, value in enumerate(values):
                key = aliases.get(_header(value))
                if key and key not in columns:
                    columns[key] = i
            if 'date' in columns and 'amount' in columns:
                headers = [_text(v) for v in values]
                break
        else:
            raise ValueError('Date and amount headers not found in first 20 rows')
        result = dict(name=Path(filename).name, sheet=selected, available_sheets=sheets,
                      headers=headers, header_row=header_row, entries=[], warnings=[], skipped=0)
        references = set()
        for offset, values in enumerate(rows, 1):
            if offset > MAX_ROWS:
                result['warnings'].append('Row limit: only first 20000 data rows imported')
                break
            row = header_row + offset
            if not any(_text(v) for v in values):
                continue
            def get(key):
                index = columns.get(key)
                return values[index] if index is not None and index < len(values) else None
            amount, day = _amount(get('amount')), _date(get('date'))
            if amount is None or day is None:
                result['skipped'] += 1
                if result['skipped'] <= 100:
                    result['warnings'].append(f'Row {row}: invalid date or amount; skipped')
                continue
            description = _text(get('description'))
            payer = _text(get('payer'))
            if not payer:
                match = re.search(r'\bUN\b\s*[:/=-]?\s*(.+)$', description, re.I)
                payer = match[1].strip() if match else ''
            ref = _text(get('ref'))
            if ref and ref in references:
                result['warnings'].append(f'Row {row}: duplicate reference {ref}')
            if ref:
                references.add(ref)
            result['entries'].append(dict(id=sha256(f'{filename}|{selected}|{row}'.encode()).hexdigest(),
                row=row, date=day, amount=amount, currency=_currency(get('currency')) or _currency(currency),
                payer=payer, ref=ref, description=description, status=_text(get('status'))))
        if result['skipped']:
            result['warnings'].append(f"Skipped {result['skipped']} invalid date/amount rows")
        return result
    finally:
        if book is not None:
            book.close()


def _business_ids(text):
    text = re.sub(r'https?://\S+|\b\S+@\S+\b', ' ', text, flags=re.I).upper()
    pattern = r'(?<![A-Z0-9])(?:SLIND-\d{4}-\d+|SLINXTF\d+|S20\d{10,20}|SIT[A-Z0-9]{7,30}|[A-Z]{4}\d{7})(?![A-Z0-9])'
    return list(dict.fromkeys(v for v in re.findall(pattern, text) if any(c.isdigit() for c in v)))


def effective_payment_date(payment: dict) -> dict:
    """Apply the date rule at read time too, so old OCR indexes remain usable."""
    if payment.get('ambiguous') or _date(payment.get('date')):
        return payment
    day = _date(payment.get('received'))
    if not day:
        return payment
    result = dict(payment, date=day, date_source='mail_received')
    evidence = dict(payment.get('evidence') or {})
    evidence['date'] = [dict(value=day, source='mail_received', quote=_text(payment.get('received')),
                             note='未识别到明确付款日期，使用邮件时间匹配，并非凭证付款日期。')]
    result['evidence'] = evidence
    note = '未识别到明确付款日期，使用邮件时间匹配。'
    result['warnings'] = list(dict.fromkeys([w for w in payment.get('warnings', []) if w not in ('Missing date', 'Missing date: date similarity receives 0 points')] + [note]))
    remark = _text(payment.get('remark')).replace('日期未识别', day + '（邮件时间代用）')
    if '邮件时间代用' not in remark:
        remark += ('\n' if remark else '') + '匹配日期：' + day + '（邮件时间代用）'
    result['remark'] = remark
    fingerprint = '|'.join((_name(result.get('payer')), day, result.get('amount') or '', result.get('currency') or ''))
    result['duplicate_key'] = sha256(fingerprint.encode()).hexdigest()
    return result


def build_payments(mail: dict) -> list[dict]:
    """One payment per attachment page; never combine parts across pages/mail.

    business_ids from mail are related clues, not confirmed allocation. Source
    evidence keeps OCR line/quote/confidence/box and its attachment identifier.
    """
    pages = []
    for asset in mail.get('attachment_details') or []:
        for page in (asset.get('result') or {}).get('pages') or []:
            fields = page.get('fields') or {}
            if any(fields.get(k) for k in ('payer', 'date', 'amount', 'currency', 'identifiers', 'fee')):
                pages.append((asset, page.get('number', 1), fields, 'attachment'))
    if not pages and _text(mail.get('latest_body')):
        from app.mail_attachments import extract_fields
        lines = [dict(text=line, confidence=None, box=None) for line in mail['latest_body'].splitlines() if line.strip()]
        fields = extract_fields(lines, 1)
        # Only explicitly date-labelled latest-body lines may supply a date.
        fields['date'] = [f for f in fields['date'] if re.match(
            r'^\s*(?:payment date|transfer date|transaction date|value date|date|日期|付款日期|转账日期)\s*[:：]', f['quote'], re.I)]
        if any(fields.get(k) for k in ('payer', 'amount', 'date')):
            pages.append((dict(id='latest_body'), 1, fields, 'latest_body'))
    result = []
    mail_clues = _business_ids('\n'.join(_text(mail.get(k)) for k in ('subject', 'latest_body', 'quoted_body')))
    for asset, page, fields, source in pages:
        warnings, values, evidence = [], {}, {}
        ambiguous = False
        for key, normalizer in (('payer', _name), ('date', _date), ('amount', _amount), ('currency', _currency)):
            items = fields.get(key) or []
            if key == 'currency':
                items = items + [dict(value=item.get('currency')) for item in fields.get('amount') or [] if item.get('currency')]
            evidence[key] = [dict(item, attachment_id=asset.get('id'), source=source) for item in items]
            distinct = {}
            for item in items:
                normalized = normalizer(item.get('value'))
                if normalized:
                    distinct[normalized] = item.get('value') if key == 'payer' else normalized
                elif _text(item.get('value')):
                    warnings.append(f'Invalid {key} field')
                    ambiguous = True
            if len(distinct) > 1:
                warnings.append(f'Ambiguous page: multiple {key} values')
                ambiguous = True
            values[key] = next(iter(distinct.values())) if len(distinct) == 1 else None
        confirmed = _business_ids(' '.join(_text(item.get('value')) for item in fields.get('identifiers') or []))
        evidence['identifiers'] = [dict(item, attachment_id=asset.get('id'), source=source) for item in fields.get('identifiers') or []]
        evidence['related_mail_clues'] = dict(business_ids=mail_clues, allocation_confirmed=False)
        ids = list(dict.fromkeys(confirmed + mail_clues))
        if mail_clues:
            warnings.append('Related mail business IDs are clues; payment allocation is not confirmed')
        if ambiguous:
            values = dict.fromkeys(('payer', 'date', 'amount', 'currency'))
        for key in ('payer', 'date', 'amount', 'currency'):
            if not values[key]:
                warnings.append(f'Missing {key}')
        mid, aid = _text(mail.get('id')), _text(asset.get('id'))
        identity = sha256(f'{mid}|{aid}|{page}'.encode()).hexdigest()
        keyword=next((v for v in ids if v.startswith(('SLIND-', 'SLINXTF', 'SIT'))), _text(mail.get('subject')))
        search = '\n'.join((f'检索关键词：{keyword}',f"主题：{_text(mail.get('subject'))}",f"发件人：{_text(mail.get('sender'))}",f"邮件时间：{_text(mail.get('received'))}"))
        remarks=[]
        if confirmed: remarks.append('凭证号码：'+'、'.join(confirmed))
        if mail_clues: remarks.append('邮件关联号码（待核对）：'+'、'.join(mail_clues))
        if not ids: remarks.append('邮件主题：'+_text(mail.get('subject')))
        remarks.append(f"付款：{values['date'] or '日期未识别'}；{values['currency'] or '币种未识别'} {values['amount'] or '金额未识别'}；{values['payer'] or '付款人未识别'}")
        fingerprint = '|'.join((_name(values['payer']), values['date'] or '', values['amount'] or '', values['currency'] or ''))
        result.append(dict(id=identity, mail_id=mid, attachment_id=aid, attachment_name=asset.get('name','最新正文' if source=='latest_body' else '付款凭证'), page=page, **values,
            subject=_text(mail.get('subject')), sender=_text(mail.get('sender')), received=_text(mail.get('received')),
            keyword=keyword,
            business_ids=ids, remark='\n'.join(remarks), search_info=search, evidence=evidence, warnings=warnings,
            duplicate_key=sha256(fingerprint.encode()).hexdigest(), attachment_sha256=asset.get('sha256') or '',
            ambiguous=ambiguous))
    return [effective_payment_date(payment) for payment in result]


def _payer_score(left, right):
    left, right = _name(left), _name(right)
    if not left or not right:
        return 0, 'Missing payer'
    if left == right:
        return 30, None
    if min(len(left), len(right)) >= 8 and (left.startswith(right) or right.startswith(left)):
        return 24, 'Payer prefix truncation'
    ratio = SequenceMatcher(None, left, right).ratio()
    if ratio >= .8:
        return min(23, round(23 * ratio)), 'Fuzzy payer match'
    return 0, 'Payer conflict'


def _group_payments(payments):
    """Group complete evidence only; same bytes use page in their identity."""
    groups, indexes = [], {}
    def signature(payment):
        return (_name(payment.get('payer')), _date(payment.get('date')),
                _amount(payment.get('amount')), _currency(payment.get('currency')),
                bool(payment.get('ambiguous')))
    for payment in payments:
        keys = []
        if payment.get('attachment_sha256'):
            keys.append(('sha', payment['attachment_sha256'], payment.get('page')))
        complete = all(payment.get(k) is not None and payment.get(k) != '' for k in ('payer', 'date', 'amount', 'currency'))
        if complete and not payment.get('ambiguous'):
            # Only an invoice ID is strong enough to group different source bytes.
            invoices = sorted(v for v in payment.get('business_ids', []) if re.fullmatch(r'SLIND-\d{4}-\d+|SLINXTF\d+', v))
            if invoices:
                keys.append(('fingerprint', _name(payment['payer']), _date(payment['date']), _amount(payment['amount']),
                             _currency(payment['currency']), tuple(invoices)))
        found = next((indexes[k] for k in keys if k in indexes), None)
        # Do not hide differing interpretations of the same source page.
        if found is not None and signature(groups[found][0]) != signature(payment):
            found = None
        if found is None:
            found = len(groups)
            groups.append([])
        groups[found].append(payment)
        for key in keys:
            indexes.setdefault(key, found)
    return groups


def rank_receipts(receipts: list, payments: list, day_window=3, limit=3, mapping=None, payment_id=None) -> list:
    """Date window is mandatory; amount/currency differences lower similarity.

    candidate_count is the deduplicated count before limit. Related search info
    is a list of {mail_id, search_info}; no scores are added for repeated mail.
    """
    if day_window < 0 or limit < 0:
        raise ValueError('day_window and limit must be nonnegative')
    groups = _group_payments([effective_payment_date(payment) for payment in payments])
    if payment_id is not None:
        # A saved confirmation may reference a noncanonical source copy. Keep
        # its original provenance, while retaining all copies for mapping clues.
        groups=[[selected]+[p for p in group if p is not selected]
                for group in groups for selected in group if selected.get('id')==payment_id]
    by_date={}
    for group in groups:
        day=_date(group[0].get('date'))
        if day and not group[0].get('ambiguous'):
            by_date.setdefault(day,[]).append(group)
    results = []
    for receipt in receipts:
        candidates = []
        amount = _amount(receipt.get('amount'))
        receipt_day=_date(receipt.get('date'))
        nearby=[]
        if receipt_day:
            day=date.fromisoformat(receipt_day)
            for offset in range(-day_window,day_window+1):
                nearby.extend(by_date.get((day+timedelta(days=offset)).isoformat(),[]))
        for group in nearby if amount is not None else []:
            payment = group[0]
            other_amount=_amount(payment.get('amount'))
            if payment.get('ambiguous') or other_amount is None:
                continue
            payer_score, warning = _payer_score(receipt.get('payer'), payment.get('payer'))
            if other_amount!=amount and payer_score<18:
                continue
            rc, pc = _currency(receipt.get('currency')), _currency(payment.get('currency'))
            warnings = list(payment.get('warnings') or [])
            if other_amount==amount:
                amount_score=50
            else:
                delta=abs(Decimal(amount)-Decimal(other_amount))
                base=max(abs(Decimal(amount)),abs(Decimal(other_amount)))
                relative=delta/base if base else Decimal(0)
                amount_score=35 if relative<=Decimal('.01') else 25 if relative<=Decimal('.05') else 15 if relative<=Decimal('.10') else 0
                warnings.append(f'金额不一致：流水 {amount}，邮件 {other_amount}；请核对差额或拆分付款。')
            if rc and pc and rc!=pc:
                amount_score=0
                warnings.append(f'币种不一致：流水 {rc}，邮件 {pc}；金额项计 0 分，未进行汇率换算，请核对原始凭证。')
            elif not rc or not pc:
                amount_score=0
                warnings.append('Unknown currency: amount similarity receives 0 points; no currency inferred')
            rd, pd = _date(receipt.get('date')), _date(payment.get('date'))
            date_score = 0
            if rd and pd:
                gap = abs((date.fromisoformat(rd) - date.fromisoformat(pd)).days)
                if gap > day_window:
                    continue
                date_score = max(0, 20 - gap * 5)
            else:
                # An undated candidate must never bypass the configured window.
                continue
            if warning:
                warnings.append(warning.lower())
            breakdown = dict(amount=amount_score, payer=payer_score, date=date_score)
            relationship=None
            if (mapping is not None and payer_score<30 and amount==other_amount
                    and rc in MAPPING_CURRENCIES and rc==pc):
                relationship=mapping.match(receipt.get('payer'),group)
                if relationship and relationship['points']>payer_score:
                    breakdown['mapping']=relationship['points']-payer_score
                    warnings.append('客户台账包含共享邮箱或一对多客户关系，请逐项核对。' if relationship['ambiguous'] else '通过客户台账关系补充匹配；未替换凭证付款人，请核对原始凭证。')
                else:relationship=None
            candidate = dict(payment, payment_id=payment.get('id'), score=sum(breakdown.values()), breakdown=breakdown,
                warnings=list(dict.fromkeys(warnings)), related_mail_ids=list(dict.fromkeys(p.get('mail_id') for p in group)),
                related_search_info=[dict(mail_id=p.get('mail_id'), search_info=p.get('search_info', '')) for p in group])
            if relationship:
                from copy import deepcopy
                candidate['mapping']=deepcopy(relationship)
                candidate['evidence']=deepcopy(candidate.get('evidence') or {})
                candidate['evidence']['customer_mapping']=deepcopy(relationship)
                lines=[f"{row['bank_name']} → {row['customer']}；邮箱：{', '.join(row['emails']) or '未登记'}；{row['sheet']} 第{row['row']}行" for row in relationship['rows']]
                candidate['remark']=str(candidate.get('remark') or '')+'\n客户台账关联（辅助依据）：'+'；'.join(lines)
            candidates.append(candidate)
        candidates.sort(key=lambda c: (-c['score'], _text(c.get('id'))))
        results.append(dict(receipt=receipt, candidates=candidates[:limit], candidate_count=len(candidates),
                            warnings=['Scores indicate similarity, not probability']))
    return results
