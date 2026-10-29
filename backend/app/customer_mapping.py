"""Explicit customer/payer/email relationships; never modify original evidence."""
from collections import defaultdict
import csv
from email.utils import getaddresses
from hashlib import sha256
from io import BytesIO, StringIO
import json
from pathlib import Path
import re
from zipfile import ZipFile

from .mail_matching import _name

MAX_BYTES=5*1024**2
MAX_ROWS=20000
HEADERS={
    'customer':('Customer Name','客户名称','客户公司','客户'),
    'bank_name':('Customer Bank Name Reff','Customer Bank Name Ref','Customer Bank Name','银行付款人','付款人','银行户名'),
    'email':('Email','E-mail','邮箱','客户邮箱'),
}


def emails(value):
    text=str(value or '').replace(';',',').replace('；',',').replace('\n',',')
    return list(dict.fromkeys(address.casefold() for _,address in getaddresses([text])
                             if re.fullmatch(r'[^\s@<>]+@[^\s@<>]+\.[^\s@<>]+',address)))


def header(value):return re.sub(r'[^\w]','',str(value or '').casefold())


def parse_contacts(raw,name,sheet_name=None):
    if not raw or len(raw)>MAX_BYTES:raise ValueError('客户台账不能为空，且不能超过5MB。')
    suffix=Path(name).suffix.lower();book=None
    try:
        if suffix=='.xlsx':
            from openpyxl import load_workbook
            with ZipFile(BytesIO(raw)) as archive:
                if sum(item.file_size for item in archive.infolist())>100*1024**2:
                    raise ValueError('客户台账解压后过大，请拆分文件。')
            book=load_workbook(BytesIO(raw),read_only=True,data_only=True)
            sheets=book.sheetnames
            if sheet_name and sheet_name not in sheets:raise ValueError('找不到指定的客户台账工作表。')
            selected=sheet_name or sheets[0]
            rows=book[selected].iter_rows(max_col=30,values_only=True)
        elif suffix=='.csv':
            sheets=['CSV'];selected='CSV';text=None
            for encoding in ('utf-8-sig','gb18030'):
                try:text=raw.decode(encoding);break
                except UnicodeDecodeError:pass
            if text is None:raise ValueError('客户台账编码无法识别，请导出UTF-8 CSV或XLSX。')
            try:dialect=csv.Sniffer().sniff(text[:8192],delimiters=',;\t')
            except csv.Error:dialect=csv.excel
            rows=csv.reader(StringIO(text),dialect)
        else:raise ValueError('请选择XLSX或CSV客户台账文件。')
        columns=None;entries=[];excluded=[];skipped=0;duplicates=0;seen={}
        aliases={header(alias):key for key,values in HEADERS.items() for alias in values}
        for number,row in enumerate(rows,1):
            if number>MAX_ROWS+20:raise ValueError('客户台账超过20000行，请拆分文件。')
            if columns is None:
                found={aliases[header(value)]:index for index,value in enumerate(row) if header(value) in aliases}
                if 'customer' in found and 'bank_name' in found:columns=found
                elif number>=20:raise ValueError('未找到客户名称和银行付款人列，请检查表头。')
                continue
            if not any(str(v or '').strip() for v in row):continue
            def value(key):
                index=columns.get(key)
                return str(row[index] or '').strip() if index is not None and index<len(row) else ''
            customer,bank,email=value('customer'),value('bank_name'),value('email')
            addresses=emails(email)
            reasons=[]
            if not customer:reasons.append('客户名称缺失')
            if not bank:reasons.append('银行付款人缺失')
            if len(customer)>300:reasons.append('客户名称超过300字')
            if len(bank)>300:reasons.append('银行付款人超过300字')
            if email and not addresses:reasons.append('邮箱格式无法识别')
            source=dict(customer=customer[:1000],bank_name=bank[:1000],email=email[:1000],row=number,sheet=selected,
                        truncated=any(len(v)>1000 for v in (customer,bank,email)))
            if reasons:
                excluded.append(dict(source,reasons=reasons,kind='invalid'))
                skipped+=1;continue
            signature=(customer,bank,tuple(addresses))
            if signature in seen:
                duplicates+=1
                excluded.append(dict(source,reasons=['与已有映射完全重复，已合并'],kind='duplicate',duplicate_of=seen[signature]))
                continue
            seen[signature]=number
            entries.append(dict(customer=customer,bank_name=bank,emails=addresses,row=number,sheet=selected))
        if not columns or not entries:raise ValueError('客户台账没有有效映射；每行至少需要客户名称和银行付款人，邮箱可留空。')
        revision=sha256(json.dumps(entries,ensure_ascii=False,sort_keys=True).encode()).hexdigest()
        return dict(name=name,sheet=selected,available_sheets=sheets,entries=entries,entries_count=len(entries),skipped=skipped,duplicates_count=duplicates,excluded_rows=excluded,revision=revision,enabled=True)
    except ValueError:raise
    except Exception as error:raise ValueError('客户台账无法解析，请检查文件格式或重新导出。') from error
    finally:
        if book:book.close()


class MappingIndex:
    def __init__(self,entries):
        self.by_payer=defaultdict(list);self.by_email=defaultdict(set)
        for row in entries:
            for key in {_name(row['bank_name']),_name(row['customer'])}:
                if key:self.by_payer[key].append(row)
            for address in row['emails']:self.by_email[address].add((_name(row['customer']),_name(row['bank_name'])))

    def match(self,payer,group):
        rows=self.by_payer.get(_name(payer),[])
        if not rows:return None
        shared_payer=len({(_name(row['customer']),_name(row['bank_name'])) for row in rows})>1
        matches=[];points=0;ambiguous=False
        for payment in group:
            name=_name(payment.get('payer'));senders=set(emails(payment.get('sender')))
            for row in rows:
                aliases={_name(row['customer']),_name(row['bank_name'])}
                # A concrete, unrelated payer must not be overridden by email.
                if name and name not in aliases:continue
                matched=senders.intersection(row['emails'])
                if not name and not matched:continue
                shared=shared_payer or any(len(self.by_email[address])>1 for address in matched)
                score=10 if shared else 22 if name else 24
                if row not in matches:matches.append(row)
                points=max(points,score);ambiguous=ambiguous or shared
        if not matches:return None
        # Any ambiguous path lowers the whole relationship, not just one row.
        return dict(points=10 if ambiguous else points,ambiguous=ambiguous,rows=matches,
                    label='客户台账映射（共享关系，需核对）' if ambiguous else '客户台账映射')


def preview_contacts(snapshot,status='valid',keyword='',page=1,page_size=50,search_aliases=()):
    """Read an encrypted-store snapshot only; do not reparse or fetch a source."""
    if status not in ('valid','excluded') or page<1 or not 1<=page_size<=100:
        raise ValueError('预览筛选或页码无效。')
    metadata={k:v for k,v in snapshot.items() if k not in ('entries','excluded_rows')}
    diagnostics_available='excluded_rows' in snapshot
    notice=''
    if not diagnostics_available and snapshot.get('skipped',0):
        notice='此快照由旧版本读取，仅保存未导入数量，未保存具体行和原因。请手动重新读取客户台账后查看；打开预览不会触发读取。'
    entries=snapshot.get('entries',[])
    index=MappingIndex(entries) if status=='valid' else None
    multiple={key for key,rows in index.by_payer.items() if len({(_name(item['customer']),_name(item['bank_name'])) for item in rows})>1} if index else set()
    source=entries if status=='valid' else snapshot.get('excluded_rows',[])
    needle=keyword.strip().casefold();needles=[needle]+[alias.casefold() for alias in search_aliases];matched=[];total=0
    start=(page-1)*page_size
    for row in source:
        flags=[]
        if index is not None:
            if not row['emails']:flags.append('未登记邮箱')
            if any(len(index.by_email[address])>1 for address in row['emails']):flags.append('共享邮箱')
            if any(key in multiple for key in {_name(row['customer']),_name(row['bank_name'])}):flags.append('一对多关系')
        searchable=[row.get('customer',''),row.get('bank_name',''),row.get('email',''),*row.get('emails',[]),
                    str(row.get('row','')),row.get('sheet',''),*row.get('reasons',[]),*flags]
        if needle and not any(term in ' '.join(str(v) for v in searchable).casefold() for term in needles):continue
        if start<=total<start+page_size:
            matched.append({**row,'flags':flags})
        total+=1
    return dict(metadata=metadata,rows=matched,total=total,page=page,page_size=page_size,
                diagnostics_available=diagnostics_available,notice=notice)
