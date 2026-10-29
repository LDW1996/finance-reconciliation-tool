"""Response-only localization. Original business strings are never traversed as text."""
import re

EXACT = {'空闲':'Idle','读取中':'Reading','已停止':'Stopped','完成':'Complete','失败':'Failed','等待中':'Waiting','扫描中':'Scanning','邮件记录不存在，请刷新列表。':'Mail record not found. Refresh the list.','附件不存在。':'Attachment not found.','页码无效。':'Invalid page number.','包含未知的保存设置。':'Unknown storage setting.','请先确认存储优化。':'Confirm storage optimization first.'}
EXACT.update({'分数表示规则相似度，不代表正确概率。': 'Scores indicate similarity, not probability', '邮件中的业务号码为关联线索，请核对是否对应本笔付款。': 'Related mail business IDs are clues; payment allocation is not confirmed', '付款人缺失，付款人项计 0 分。': 'missing payer', '付款人缺失，请查看凭证。': 'Missing payer', '付款日期缺失，请查看凭证。': 'Missing date', '付款金额缺失，请查看凭证。': 'Missing amount', '币种缺失，请查看凭证。': 'Missing currency', '付款人不一致，请重点核对。': 'payer conflict', '银行付款人名称存在前缀截断，付款人项计 24 分。': 'payer prefix truncation', '付款人名称近似，请核对原文。': 'fuzzy payer match', '币种未明确识别，金额项计 0 分。': 'Unknown currency: amount similarity receives 0 points; no currency inferred', '付款日期缺失，日期项计 0 分。': 'Missing date: date similarity receives 0 points'})


EXACT.update({
 '客户名称缺失':'Customer name is missing','银行付款人缺失':'Bank payer is missing',
 '客户名称超过300字':'Customer name exceeds 300 characters','银行付款人超过300字':'Bank payer exceeds 300 characters',
 '邮箱格式无法识别':'Email address format is invalid','与已有映射完全重复，已合并':'Identical mapping merged',
 '未登记邮箱':'No email registered','共享邮箱':'Shared email','一对多关系':'One-to-many relationship',
 '客户台账映射':'Customer ledger mapping','客户台账映射（共享关系，需核对）':'Customer ledger mapping (shared relationship; review required)',
 '客户台账包含共享邮箱或一对多客户关系，请逐项核对。':'Customer ledger contains shared email or one-to-many relationships. Review each match.',
 '通过客户台账关系补充匹配；未替换凭证付款人，请核对原始凭证。':'Customer ledger relationship supports this match. The receipt payer is unchanged; check the original receipt.',
 '此快照由旧版本读取，仅保存未导入数量，未保存具体行和原因。请手动重新读取客户台账后查看；打开预览不会触发读取。':'This older snapshot contains excluded row counts only. Refresh the customer ledger manually to see row details and reasons. Opening preview does not refresh it.',
 '磁盘剩余空间不足，已暂停新增读取；请释放空间后重试。':'Low disk space. New mail reads are paused; free disk space and retry.',
 '邮件存储达到容量上限，已暂停新增读取；请调整容量或释放空间后重试。':'Mail storage limit reached. New reads are paused; increase the limit or free space and retry.',
 '去重和预览清理已完成，但数据库整理所需临时空间不足；请释放空间后再次优化。':'Deduplication and preview cleanup completed. Database compaction needs more temporary disk space; free space and optimize again.',
 '分析进程意外退出':'Analysis worker exited unexpectedly','重新建立分析进程后重试':'Restart the analysis worker and retry',
 '数据库暂时被占用':'Database is temporarily busy','稍后重新执行任务':'Retry the task shortly',
 '存储空间不足':'Insufficient disk space','释放磁盘空间后手动重试':'Free disk space and retry manually',
 '数据库文件异常':'Database file is damaged','检查并修复数据库或从备份恢复':'Check and repair the database or restore a backup',
 '存储访问权限不足':'Storage access denied','检查存储访问权限后手动重试':'Check storage permissions and retry manually',
 '可用内存不足':'Insufficient available memory','释放内存后手动重试':'Free memory and retry manually',
 '存储读写失败':'Storage read/write failed','检查磁盘和存储连接后手动重试':'Check the disk and storage connection, then retry manually',
 '数据库操作失败':'Database operation failed','检查数据库状态后手动重试':'Check database status and retry manually',
 '存储路径或文件不存在':'Storage path or file is missing','检查数据目录是否被移动、删除或断开；恢复原目录后手动重试':'Check whether the data directory was moved, deleted or disconnected. Restore the original directory and retry manually',
 '文件读写失败':'File read/write failed','检查存储状态后手动重试':'Check storage status and retry manually',
 '任务执行异常':'Task failed unexpectedly','检查任务状态后手动重试':'Check task status and retry manually',
 '列出 Outlook 已打开的邮箱和 PST 存档。未打开的存档不在列表中。':'Lists mailboxes and PST archives currently open in Outlook. Closed archives are excluded.',
 '上次存储优化被中断；数据保留，可重新预览并继续。':'Previous storage optimization was interrupted. Data is preserved; preview again to continue.',
 '任务已恢复，正在继续处理后续邮件。':'Task recovered. Processing subsequent mail.',
})
EXACT.update({
 '等待识别的邮件较多，正在消化队列；积压下降后自动读取将提前继续。':'Recognition backlog is high. Processing the queue; automatic reading will resume when the backlog falls.',
 '附件处理已因明确故障暂停，新增读取暂缓。请按上方故障提示处理后手动重试。':'Attachment processing is paused due to a fault. Resolve the diagnostic above and retry manually.',
 '未识别到明确付款日期，使用邮件时间匹配。':'No explicit payment date was recognized; the mail received date is used for matching.',
 '未识别到明确付款日期，使用邮件时间匹配，并非凭证付款日期。':'The mail received date is used for matching; it is not a recognized payment date.',
 '来源邮件或匹配条件已更新，请重新核对关联。':'Source mail or matching conditions changed. Review this association again.',
 '付款证据已更新，请刷新后重新核对。':'Payment evidence changed. Refresh and review again.',
 '请先配置客户台账的在线来源和授权。':'Configure the online customer source and authorization first.',
 '请先配置银行流水授权，或选择独立授权并上传服务账号JSON。':'Configure bank authorization first, or select separate authorization and upload a service-account JSON file.',
 '授权文件无效，请选择 Google 服务账号的 JSON 密钥文件。':'Invalid authorization file. Select a Google service-account JSON key.',
 '无读取权限，请启用 Google Sheets API，并将客户台账以查看权限共享给服务账号。':'Read access denied. Enable the Google Sheets API and share the customer sheet with the service account as a viewer.',
 '没有已确认的邮件关联，请先核对并确认。':'No confirmed mail associations. Review and confirm an association first.',
 '确认内容已变化，请刷新后重新确认回写。':'Confirmed content changed. Refresh and confirm writeback again.',
 '已确认关联数量已变化，请刷新后重新确认回写。':'The confirmed association count changed. Refresh and confirm writeback again.',
})
EXACT.update(dict(zip(
 ['任务调度','读取状态','恢复队列断点','建立复用索引','启动识别进程','检查存储空间','领取识别任务','记录并跳过异常邮件','保存识别队列进度','恢复自动读取','保存邮件与读取进度','读取 Outlook','准备检查存储','等待当前附件识别完成，剩余队列已暂停','复查所选邮件并保存加密恢复备份','恢复清理备份，保留后来读取的数据','检查完成','恢复完成','清理完成','存储优化完成','查询邮件存储占用','检查历史附件','去重并清理可重建预览','整理数据库以释放磁盘空间','存储维护'],
 ['Scheduling task','Reading status','Restoring queue checkpoint','Building reuse index','Starting recognition worker','Checking storage space','Claiming recognition task','Recording and skipping failed mail','Saving recognition queue progress','Resuming automatic reads','Saving mail and read progress','Reading Outlook','Preparing storage check','Waiting for current attachment recognition; remaining queue paused','Reviewing selected mail and saving encrypted recovery backup','Restoring cleanup backup while preserving later data','Check complete','Restore complete','Cleanup complete','Storage optimization complete','Checking mail storage usage','Checking historical attachments','Deduplicating and cleaning rebuildable previews','Compacting database to release disk space','Storage maintenance'])))

# These keys contain application-generated diagnostics, not imported business text.
GENERATED={'message','detail','error','errors','warnings','warning','phase','sync_phase','note','notes',
 'status_text','progress_text','analysis_error','analysis_recovery','reason','reasons','action',
 'flags','notice','blocked_reason','compact_note'}
LABELS={'凭证号码':'Receipt reference','邮件关联号码（待核对）':'Related mail references (review required)',
 '邮件主题':'Mail subject','匹配日期':'Matching date','客户台账关联（辅助依据）':'Customer ledger association (supporting evidence)',
 '发件人':'Sender','接收时间':'Received','收件时间':'Received','邮件时间':'Mail received','主题':'Subject','附件':'Attachment',
 '付款人':'Payer','日期':'Date','金额':'Amount','币种':'Currency','邮件':'Mail','关键词':'Keyword','检索关键词':'Search keyword','客户':'Customer'}

def text(value):
    """Translate generated diagnostics only; unknown text gets a stable support code."""
    if not isinstance(value,str) or not re.search('[\u4e00-\u9fff]',value):return value
    if value in EXACT:return EXACT[value]
    value=re.sub(r'第 (\d+) 行日期或金额无效，已跳过。',r'Row \1: invalid date or amount; skipped',value)
    value=re.sub(r'共跳过 (\d+) 行日期或金额无效的记录。',r'Skipped \1 rows with invalid date or amount.',value)
    # Diagnostic composition contains only curated reason/stage/action strings.
    for source,target in sorted(EXACT.items(),key=lambda item:-len(item[0])):
        value=value.replace(source,target)
    value=re.sub(r'自动恢复第 (\d+)/(\d+) 次',r'Automatic recovery attempt \1/\2',value)
    value=value.replace('已达到恢复上限，任务暂停','Recovery limit reached; task paused').replace('任务暂停，需处理后手动重试','Task paused; resolve the issue and retry manually')
    if not re.search('[\u4e00-\u9fff]',value):return value
    import hashlib
    return 'Operation requires attention (diagnostic '+hashlib.sha256(value.encode()).hexdigest()[:10]+').'

def generated_lines(value):
    """Parse known generated grammar; never replace inside business value fields."""
    lines=[]
    for original in value.split('\n'):
        line=original
        prefix=re.match(r'^(已确认关联|待核对候选)( \d+)?：(.*)$',line)
        if prefix:
            label='Confirmed association' if prefix[1]=='已确认关联' else 'Candidate for review'
            lines.append(label+(prefix[2] or '')+':'+generated_lines(prefix[3]));continue
        candidate=re.fullmatch(r'候选 (\d+)（(\d+) 分）',line)
        if candidate:lines.append(f'Candidate {candidate[1]} (score {candidate[2]})');continue
        if line in {'暂不关联','未匹配到候选邮件'}:
            lines.append({'暂不关联':'Association rejected','未匹配到候选邮件':'No candidate mail matched'}[line]);continue
        if line.startswith('付款：'):
            # The writer emits date; currency amount; payer. Payer stays verbatim.
            parts=line[3:].split('；',2)
            if len(parts)==3:
                date,currency_amount,payer=parts
                date='Date unavailable' if date=='日期未识别' else date
                currency,sep,amount=currency_amount.partition(' ')
                currency='Currency unavailable' if currency=='币种未识别' else currency
                amount='Amount unavailable' if amount=='金额未识别' else amount
                payer='Payer unavailable' if payer=='付款人未识别' else payer
                line=f'Payment: Date: {date}; Currency / amount: {currency}{sep}{amount}; Payer: {payer}'
                lines.append(line);continue
        match=re.match(r'^(.*?)[：:](.*)$',line)
        if match and match[1] in LABELS:
            content=match[2]
            if match[1]=='客户台账关联（辅助依据）':
                # Mapping writer grammar: bank → customer; email:...; sheet row N.
                content=re.sub(r'；邮箱：([^；]*)；([^；]*) 第(\d+)行',lambda m:'; Email: '+('Not registered' if m[1]=='未登记' else m[1])+'; '+m[2]+' row '+m[3],content)
            elif match[1]=='匹配日期':
                content=re.sub(r'（邮件时间代用）$',' (mail received date used)',content)
            line=LABELS[match[1]]+':'+content
        lines.append(line)
    return '\n'.join(lines)

def response(value, key=None, *, _path=()):
    """Localize response copies by field path. Never persist translated values."""
    path=_path+((key,) if key is not None else ())
    if isinstance(value,dict):
        result={k:response(v,k,_path=path) for k,v in value.items()}
        if {'kind','phase','reason','action','message'}<=value.keys():
            result['message']=f"{result['phase']}: {result['reason']}. {result['action']}"
            if value.get('retryable'):
                result['message']+=f" (Recovery attempt {value.get('attempt',0)}/{value.get('limit',0)})"
        diagnostic=result.get('analysis_diagnostic')
        if isinstance(diagnostic,dict):
            if value.get('analysis_error'):result['analysis_error']=diagnostic.get('message',result['analysis_error'])
            if value.get('analysis_recovery'):result['analysis_recovery']=diagnostic.get('message',result['analysis_recovery'])
        return result
    if isinstance(value,list):return [response(v,_path=path) for v in value]
    field=path[-1] if path else None
    if field in {'remark','search_info'} and isinstance(value,str):return generated_lines(value)
    if field=='label':
        return text(value) if len(path)>1 and path[-2] in {'mapping','customer_mapping'} else value
    return text(value) if field in GENERATED else value

def english_export(raw):
    """Edit only the three appended columns, never original bank cells."""
    from io import BytesIO
    from openpyxl import load_workbook
    book=load_workbook(BytesIO(raw))
    titles=['Mail association notes','Match score','Mail search information']
    for sheet in book:
        start=sheet.max_column-2
        if start<1:continue
        for row in sheet:
            if [row[start-1+i].value for i in range(3)]==['邮件关联备注','匹配分数','邮件检索信息']:
                header=row[0].row
                for i,title in enumerate(titles):sheet.cell(header,start+i,title)
                for cells in sheet.iter_rows(min_row=header+1,min_col=start,max_col=start+2):
                    for cell in cells:
                        if isinstance(cell.value,str):
                            value=generated_lines(cell.value)
                            value=re.sub(r'^(已确认关联|待核对候选)( \d+)?：',lambda m:('Confirmed association' if m[1]=='已确认关联' else 'Candidate for review')+(m[2] or '')+':',value,flags=re.M)
                            value={'暂不关联':'Association rejected','未匹配到候选邮件':'No candidate mail matched'}.get(value,value)
                            value=re.sub(r'^候选 (\d+)（(\d+) 分）',r'Candidate \1 (score \2)',value,flags=re.M)
                            cell.value=value;cell.data_type='s'
                break
    output=BytesIO();book.save(output);return output.getvalue()
