"""Local receipt workspace: encrypted import, candidate review and copy export."""

from copy import copy, deepcopy
from datetime import datetime, timezone
import hashlib
from io import BytesIO
import json
import re
from threading import RLock


def eligible(receipt, scope):
    status=str(receipt.get('status','')).strip().upper()
    return scope=='all' or (scope=='unidentified' and status=='UNIDENTIFIED') or (scope=='unclosed' and status in ('UNCLOSED','PARTIAL CLOSED','STORAGE UNCLOSED'))


def signature(candidate):
    facts={k:candidate.get(k) for k in ('mail_id','attachment_id','page','payer','date','amount','currency','remark','search_info','evidence')}
    return hashlib.sha256(json.dumps(facts,ensure_ascii=False,sort_keys=True).encode()).hexdigest()


def translated(message):
    exact={
        'Scores indicate similarity, not probability':'分数表示规则相似度，不代表正确概率。',
        'Related mail business IDs are clues; payment allocation is not confirmed':'邮件中的业务号码为关联线索，请核对是否对应本笔付款。',
        'missing payer':'付款人缺失，付款人项计 0 分。',
        'Missing payer':'付款人缺失，请查看凭证。',
        'Missing date':'付款日期缺失，请查看凭证。',
        'Missing amount':'付款金额缺失，请查看凭证。',
        'Missing currency':'币种缺失，请查看凭证。',
        'payer conflict':'付款人不一致，请重点核对。',
        'payer prefix truncation':'银行付款人名称存在前缀截断，付款人项计 24 分。',
        'fuzzy payer match':'付款人名称近似，请核对原文。',
        'Unknown currency: amount similarity receives 0 points; no currency inferred':'币种未明确识别，金额项计 0 分。',
        'Missing date: date similarity receives 0 points':'付款日期缺失，日期项计 0 分。',
    }
    if message in exact: return exact[message]
    for key,label in [('payer','付款人'),('date','日期'),('amount','金额'),('currency','币种')]:
        message=message.replace(f'Ambiguous page: multiple {key} values',f'同页存在多个{label}，需人工核对。').replace(f'Invalid {key} field',f'{label}无法可靠识别。')
    message=re.sub(r'Row (\d+): invalid date or amount; skipped',r'第 \1 行日期或金额无效，已跳过。',message)
    message=re.sub(r'Row (\d+): duplicate reference (.+)',r'第 \1 行存在重复参考号：\2',message)
    message=re.sub(r'Skipped (\d+) invalid date/amount rows',r'共跳过 \1 行日期或金额无效的记录。',message)
    return message.replace('Row limit: only first 20000 data rows imported','仅导入前 20,000 行流水，请拆分文件后继续。')


class MailWorkspace:
    def __init__(self,store):
        self.store=store;self.lock=RLock();self.cache=None
        with store.db() as db:
            db.execute('CREATE TABLE IF NOT EXISTS mail_ledger (id TEXT PRIMARY KEY, metadata TEXT NOT NULL, content BLOB NOT NULL)')
            db.execute('CREATE TABLE IF NOT EXISTS mail_review (receipt_id TEXT PRIMARY KEY, payload TEXT NOT NULL)')

    def _dataset(self):
        with self.store.db() as db:
            row=db.execute('SELECT id,metadata FROM mail_ledger LIMIT 1').fetchone()
            return self.store.decode(row[1]) if row else None

    def ledger(self):
        data=self._dataset()
        return {k:v for k,v in data.items() if k not in ('entries','headers')} if data else None

    def import_ledger(self,raw,name,sheet_name=None,currency='IDR'):
        from .mail_matching import parse_ledger
        try: parsed=parse_ledger(raw,name,sheet_name,currency)
        except ValueError as error:
            errors={'Ledger exceeds 25 MB':'流水文件不能超过 25 MB。','Unknown ledger sheet':'找不到指定工作表，请核对工作表名称。',
                    'Date and amount headers not found in first 20 rows':'前 20 行未找到日期和收款金额表头，请检查文件。'}
            raise ValueError(errors.get(str(error),str(error))) from None
        except Exception as error: raise ValueError('流水文件无法解析，请检查文件格式、密码或损坏情况。') from error
        parsed['sheet']=parsed['sheet'] or 'CSV'
        parsed['warnings']=[translated(w) for w in parsed.get('warnings',[])]
        identity=hashlib.sha256(raw+str(parsed['sheet']).encode()+currency.encode()).hexdigest()
        data={**parsed,'id':identity,'entries_count':len(parsed['entries']),'imported_at':datetime.now(timezone.utc).isoformat()}
        with self.lock,self.store.db() as db:
            previous=db.execute('SELECT id FROM mail_ledger LIMIT 1').fetchone()
            if not previous or previous[0]!=identity:
                db.execute('DELETE FROM mail_review')
            db.execute('DELETE FROM mail_ledger')
            db.execute('INSERT INTO mail_ledger VALUES(?,?,?)',(identity,self.store.encode(data),self.store.cipher.encrypt(raw)))
            self.cache=None
        return self.ledger()

    def _review_map(self):
        with self.store.db() as db:
            return {key:self.store.decode(payload) for key,payload in db.execute('SELECT receipt_id,payload FROM mail_review')}

    def _ranked(self,data,scope,day_window,candidate_limit):
        from .mail_matching import rank_receipts
        revision=self.store.get_state('match_revision','0')
        key=(data['id'],revision,scope,day_window,candidate_limit)
        if self.cache and self.cache[0]==key: return deepcopy(self.cache[1]),revision
        receipts=[r for r in data['entries'] if eligible(r,scope)]
        result=rank_receipts(receipts,self.store.payment_rows(),day_window=day_window,limit=candidate_limit)
        self.cache=(key,deepcopy(result))
        return result,revision

    def _candidate(self,receipt,payment_id,day_window):
        from .mail_matching import rank_receipts
        payments=[p for p in self.store.payment_rows() if p.get('id')==payment_id or p.get('payment_id')==payment_id]
        result=rank_receipts([receipt],payments,day_window=day_window,limit=1)
        return result[0]['candidates'][0] if result and result[0]['candidates'] else None

    def matching(self,scope='unidentified',page=1,page_size=50,minimum_score=0,day_window=3,candidate_limit=3,keyword='',sort_by='source'):
        with self.lock:
            data=self._dataset()
            summary=dict(total=0,eligible=0,matched=0,unmatched=0,confirmed=0)
            if not data: return dict(ledger=None,summary=summary,rows=[],revision='0',page=page,total_rows=0)
            rows,revision=self._ranked(data,scope,day_window,candidate_limit)
            reviews=self._review_map()
            summary.update(total=len(data['entries']),eligible=len(rows))
            for row in rows:
                row.setdefault('warnings',[])
                row['warnings']=[translated(w) for w in row['warnings']]
                for candidate in row['candidates']:
                    candidate['warnings']=[translated(w) for w in candidate.get('warnings',[])]
                    candidate['evidence_signature']=signature(candidate)
                review=reviews.get(row['receipt']['id'],dict(state='pending',payment_id=None))
                if review['state']=='confirmed':
                    selected=next((c for c in row['candidates'] if (c.get('payment_id') or c['id'])==review.get('payment_id')),None)
                    if selected is None: selected=self._candidate(row['receipt'],review.get('payment_id'),day_window)
                    if not selected or signature(selected)!=review.get('signature'):
                        review=dict(state='pending',payment_id=None)
                        row['warnings'].append('来源邮件或匹配条件已更新，请重新核对关联。')
                    elif all((c.get('payment_id') or c['id'])!=review['payment_id'] for c in row['candidates']):
                        selected['warnings']=[translated(w) for w in selected.get('warnings',[])]
                        selected['evidence_signature']=signature(selected)
                        row['candidates'].append(selected)
                row['review']={k:v for k,v in review.items() if k!='signature'}
                summary['confirmed']+=review['state']=='confirmed'
                summary['matched']+=bool(row['candidates'])
            summary['unmatched']=summary['eligible']-summary['matched']
            confirmed=[(r['receipt']['id'],r['review']['payment_id'],next((signature(c) for c in r['candidates'] if c['payment_id']==r['review']['payment_id']),''))
                       for r in rows if r['review']['state']=='confirmed']
            confirmation_token=hashlib.sha256(json.dumps([data['id'],day_window,sorted(confirmed)],ensure_ascii=False).encode()).hexdigest()
            if minimum_score:
                rows=[r for r in rows if r['candidates'] and max(c['score'] for c in r['candidates'])>=minimum_score]
            needle=keyword.strip().casefold()
            if needle:
                rows=[r for r in rows if any(needle in str(r['receipt'].get(k) or '').casefold() for k in ('ref','payer','description'))]
            if sort_by=='score':
                rows.sort(key=lambda r:max((c['score'] for c in r['candidates']),default=-1),reverse=True)
            total=len(rows)
            return dict(ledger=self.ledger(),summary=summary,rows=rows[(page-1)*page_size:page*page_size],
                        revision=f"{data['id']}:{revision}:{self.store.get_state('review_revision','0')}",page=page,total_rows=total,confirmation_token=confirmation_token)

    def review(self,ledger_id,receipt_id,payment_id,state,day_window,expected_signature=None):
        with self.lock:
            data=self._dataset()
            if not data or data['id']!=ledger_id: raise ValueError('流水文件已更换，请刷新后重新核对。')
            receipt=next((r for r in data['entries'] if r['id']==receipt_id),None)
            if not receipt: raise ValueError('收款记录不存在，请刷新列表。')
            result=dict(state=state,payment_id=None,updated_at=datetime.now(timezone.utc).isoformat())
            if state=='confirmed':
                candidate=self._candidate(receipt,payment_id,day_window)
                if not candidate: raise ValueError('这封邮件已不符合当前匹配条件，请刷新候选。')
                if expected_signature and signature(candidate)!=expected_signature: raise ValueError('付款证据已更新，请刷新后重新核对。')
                result.update(payment_id=payment_id,signature=signature(candidate))
            elif state not in ('pending','rejected'): raise ValueError('核对状态无效。')
            with self.store.db() as db:
                db.execute('INSERT INTO mail_review VALUES(?,?) ON CONFLICT(receipt_id) DO UPDATE SET payload=excluded.payload',(receipt_id,self.store.encode(result)))
            self.store.set_state('review_revision',result['updated_at'])
            return {'ok':True}

    def export(self,scope='unidentified',minimum_score=0,day_window=3):
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
        with self.lock:
            data=self._dataset()
            if not data: raise ValueError('请先导入银行流水。')
            result=self.matching(scope=scope,page_size=20000,minimum_score=minimum_score,day_window=day_window)
            with self.store.db() as db:
                encrypted=db.execute('SELECT content FROM mail_ledger WHERE id=?',(data['id'],)).fetchone()[0]
            raw=self.store.cipher.decrypt(encrypted)
            if data['name'].lower().endswith('.xlsx'):
                workbook=openpyxl.load_workbook(BytesIO(raw))
                sheet=workbook[data.get('online',{}).get('export_sheet',data['sheet'])]
            else:
                import pandas as pd
                workbook=openpyxl.Workbook();sheet=workbook.active;sheet.title=data['sheet'][:31]
                if data['name'].lower().endswith('.xls'):
                    frame=pd.read_excel(BytesIO(raw),sheet_name=data['sheet'],header=None,dtype=object)
                    records=frame.itertuples(index=False,name=None)
                else:
                    import csv
                    from io import StringIO
                    decoded=None
                    for encoding in ('utf-8-sig','gb18030'):
                        try: decoded=raw.decode(encoding);break
                        except UnicodeDecodeError: pass
                    if decoded is None: raise ValueError('无法重新读取流水文件编码。')
                    try: dialect=csv.Sniffer().sniff(decoded[:65536],delimiters=',;\t|')
                    except csv.Error: dialect=csv.excel
                    records=csv.reader(StringIO(decoded),dialect)
                for values in records:
                    sheet.append([None if pd.isna(v) else str(v) if not isinstance(v,(str,int,float,bool,datetime)) else v for v in values])
                # Literal source text from CSV/XLS must not become formulas on conversion.
                for row in sheet:
                    for cell in row:
                        if isinstance(cell.value,str): cell.data_type='s'
            start=sheet.max_column+1;header_row=data['header_row']
            for offset,text in enumerate(('邮件关联备注','匹配分数','邮件检索信息')):
                cell=sheet.cell(header_row,start+offset,text)
                exemplar=sheet.cell(header_row,1)
                if exemplar.has_style: cell._style=copy(exemplar._style)
                else:
                    cell.font=Font(name='Calibri',size=12,bold=True)
                    cell.fill=PatternFill('solid',fgColor='DCEAF8')
                cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
            for row in result['rows']:
                review=row['review'];candidates=row['candidates']
                confirmed=review['state']=='confirmed'
                if confirmed: candidates=[c for c in candidates if (c.get('payment_id') or c['id'])==review['payment_id']]
                if review['state']=='rejected': remarks,score,search='暂不关联',None,''
                elif not candidates: remarks,score,search='未匹配到候选邮件',None,''
                else:
                    label='已确认关联' if confirmed else '待核对候选'
                    remarks='\n\n'.join(f"{label}{'' if confirmed else f' {i}'}：{c['remark']}" for i,c in enumerate(candidates,1))
                    score=candidates[0]['score']
                    search='\n\n'.join((f"候选 {i}（{c['score']} 分）\n" if not confirmed else '')+c['search_info'] for i,c in enumerate(candidates,1))
                for offset,value in enumerate((remarks,score,search)):
                    if isinstance(value,str) and len(value)>32760: value=value[:32740]+' [内容过长，已截断]'
                    cell=sheet.cell(row['receipt']['row'],start+offset,value)
                    if isinstance(value,str): cell.data_type='s'
                    cell.font=Font(name='Calibri',size=11)
                    cell.alignment=Alignment(vertical='top',wrap_text=True)
                    if offset==1: cell.number_format='0'
                number=row['receipt']['row']
                sheet.row_dimensions[number].height=max(sheet.row_dimensions[number].height or 15,min(240,24*(max(remarks.count('\n'),search.count('\n'))+2)))
            for offset,width in enumerate((72,14,72)):
                sheet.column_dimensions[openpyxl.utils.get_column_letter(start+offset)].width=width
            output=BytesIO();workbook.save(output);workbook.close()
            return output.getvalue(),data['name'].rsplit('.',1)[0]+'_邮件匹配.xlsx'
